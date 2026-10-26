"""Testes de ponta a ponta do SisPorto (Flask test client + PostgreSQL real).

Rodar:  TEST_DATABASE_ADMIN="host=/tmp user=postgres dbname=postgres" python -m unittest -v
Cria e apaga a base 'sisporto_teste'.
"""
from __future__ import annotations

import os
import re
import unittest

from app import create_app, util
from app.db import _Conn
from app.demo import SENHA_DEMO

ADMIN_DSN = os.environ.get("TEST_DATABASE_ADMIN", "host=/tmp user=postgres dbname=postgres")
TEST_DSN = ADMIN_DSN.replace("dbname=postgres", "dbname=sisporto_teste")


def _admin_exec(sql):
    from app import db as dbmod
    if dbmod._DRIVER == "psycopg":
        import psycopg
        with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
            c.execute(sql)
    else:
        from app.pq import _lib
        conn = _lib.PQconnectdb(ADMIN_DSN.encode())
        res = _lib.PQexecParams(conn, sql.encode(), 0, None, None, None, None, 0)
        _lib.PQclear(res)
        _lib.PQfinish(conn)


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _admin_exec("DROP DATABASE IF EXISTS sisporto_teste")
        _admin_exec("CREATE DATABASE sisporto_teste")
        cls.app = create_app({"DATABASE_URL": TEST_DSN, "TESTING": True, "SECRET_KEY": "teste",
                              "SESSION_COOKIE_SECURE": False})
        with cls.app.app_context():
            from app.cli import aplicar_sql
            from app.demo import carregar
            from app.db import db
            aplicar_sql()
            cls.ids = carregar()
            db().commit()

    @classmethod
    def tearDownClass(cls):
        from flask import g  # noqa: F401
        _admin_exec("DROP DATABASE IF EXISTS sisporto_teste")

    # --- ajudantes ----------------------------------------------------------
    def cliente_http(self, email=None):
        c = self.app.test_client()
        c.get("/entrar")
        if email:
            r = self.post(c, "/entrar", {"email": email, "senha": SENHA_DEMO})
            self.assertEqual(r.status_code, 302, r.data[:300])
        return c

    def post(self, c, url, data, **kw):
        with c.session_transaction() as s:
            token = s.get("csrf")
        return c.post(url, data={**data, "csrf": token}, **kw)

    def sql(self, query, params=()):
        conn = _Conn(TEST_DSN)
        try:
            rows = conn.all(query, params)
            conn.commit()
            return rows
        finally:
            conn.close()


class TestUtil(unittest.TestCase):
    def test_duracao(self):
        p = util.parse_duracao
        self.assertEqual(p("1:30"), 90)
        self.assertEqual(p("1h30"), 90)
        self.assertEqual(p("2h"), 120)
        self.assertEqual(p("45m"), 45)
        self.assertEqual(p("45"), 45)
        self.assertEqual(p("1,5h"), 90)
        self.assertIsNone(p("abc"))

    def test_formatos(self):
        self.assertEqual(util.horas(90), "1h30")
        self.assertEqual(util.horas(120), "2h")
        self.assertEqual(util.moeda(2500), "R$ 2.500,00")
        self.assertTrue(util.cnpj_valido("11222333000181"))
        self.assertFalse(util.cnpj_valido("11222333000182"))


class TestAcesso(Base):
    def test_login_errado(self):
        c = self.cliente_http()
        r = self.post(c, "/entrar", {"email": "maria@alfalog.com.br", "senha": "errada"})
        self.assertIn("incorretos".encode(), r.data)

    def test_csrf_obrigatorio(self):
        c = self.cliente_http()
        r = c.post("/entrar", data={"email": "maria@alfalog.com.br", "senha": SENHA_DEMO})
        self.assertEqual(r.status_code, 400)

    def test_sem_login_redireciona(self):
        r = self.app.test_client().get("/chamados")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/entrar", r.headers["Location"])

    def test_cliente_nao_acessa_equipe(self):
        c = self.cliente_http("maria@alfalog.com.br")
        self.assertEqual(c.get("/chamados").status_code, 403)
        self.assertEqual(c.get("/clientes").status_code, 403)

    def test_cliente_nao_ve_chamado_de_outro(self):
        num_beta = self.sql("SELECT numero FROM chamados WHERE cliente_id = %s LIMIT 1", [self.ids["beta"]])[0]["numero"]
        c = self.cliente_http("maria@alfalog.com.br")
        self.assertEqual(c.get(f"/portal/chamados/{num_beta}").status_code, 404)

    def test_cliente_nao_ve_nota_interna(self):
        num = self.sql("SELECT numero, id FROM chamados WHERE cliente_id = %s ORDER BY numero LIMIT 1",
                       [self.ids["alfa"]])[0]
        eq = self.cliente_http("ana@portosolucoes.com.br")
        self.post(eq, f"/chamados/{num['numero']}/responder", {"texto": "SEGREDO-INTERNO", "interna": "1"})
        cli = self.cliente_http("maria@alfalog.com.br")
        r = cli.get(f"/portal/chamados/{num['numero']}")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn(b"SEGREDO-INTERNO", r.data)

    def test_tecnico_nao_administra(self):
        c = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(c.get("/equipe").status_code, 403)
        self.assertEqual(c.get("/clientes/novo").status_code, 403)


class TestTelas(Base):
    def test_telas_equipe(self):
        c = self.cliente_http("contato@portosolucoes.com.br")
        num = self.sql("SELECT numero FROM chamados LIMIT 1")[0]["numero"]
        for url in ["/painel", "/chamados", "/chamados?situacao=todos", "/chamados?busca=impressora",
                    "/chamados/novo", f"/chamados/{num}", "/clientes", f"/clientes/{self.ids['alfa']}",
                    f"/clientes/{self.ids['alfa']}/contratos/novo", f"/contratos/{self.ids['ct_alfa']}",
                    "/clientes/novo", "/equipe", "/horas", "/config/sla", "/conta/senha"]:
            with self.subTest(url=url):
                r = c.get(url)
                self.assertEqual(r.status_code, 200, f"{url}: {r.data[:500]!r}")

    def test_telas_portal(self):
        c = self.cliente_http("joao@clinicabeta.com.br")
        num = self.sql("SELECT numero FROM chamados WHERE cliente_id = %s LIMIT 1", [self.ids["beta"]])[0]["numero"]
        for url in ["/portal/", "/portal/chamados", "/portal/chamados?situacao=todos", "/portal/chamados/novo",
                    f"/portal/chamados/{num}", "/portal/horas"]:
            with self.subTest(url=url):
                r = c.get(url)
                self.assertEqual(r.status_code, 200, f"{url}: {r.data[:500]!r}")
        self.assertIn("restantes".encode(), c.get("/portal/").data)

    def test_saude(self):
        r = self.app.test_client().get("/saude")
        self.assertEqual(r.json["ok"], True)


class TestFluxo(Base):
    def test_ciclo_completo(self):
        # 1) cliente abre pelo portal
        cli = self.cliente_http("joao@clinicabeta.com.br")
        r = self.post(cli, "/portal/chamados/novo",
                      {"titulo": "Wi-Fi caiu", "descricao": "Recepção sem rede", "prioridade": "alta"})
        self.assertEqual(r.status_code, 302)
        numero = int(re.search(r"/chamados/(\d+)", r.headers["Location"]).group(1))
        ch = self.sql("SELECT * FROM chamados WHERE numero = %s", [numero])[0]
        self.assertEqual(ch["contrato_id"], self.ids["ct_beta"])        # contrato único herdado
        self.assertIsNotNone(ch["prazo_solucao"])                       # SLA calculado
        self.assertEqual(ch["origem"], "portal")

        # 2) técnico responde e pede informação ao cliente
        eq = self.cliente_http("ana@portosolucoes.com.br")
        self.post(eq, f"/chamados/{numero}/responder",
                  {"texto": "Qual o nome da rede?", "status_apos": "aguardando_cliente"})
        ch = self.sql("SELECT * FROM chamados WHERE numero = %s", [numero])[0]
        self.assertEqual(ch["status"], "aguardando_cliente")
        self.assertIsNotNone(ch["primeira_resposta_em"])
        self.assertEqual(ch["responsavel_id"], self.ids["tec"])       # quem respondeu assume

        # 3) cliente responde -> volta para atendimento
        self.post(cli, f"/portal/chamados/{numero}/responder", {"texto": "Rede RECEPCAO"})
        self.assertEqual(self.sql("SELECT status FROM chamados WHERE numero = %s", [numero])[0]["status"],
                         "em_atendimento")

        # 4) técnico lança 1h30 e resolve
        import datetime as dt
        hoje = dt.datetime.now(util.TZ).date().isoformat()
        r = self.post(eq, f"/chamados/{numero}/horas",
                      {"data": hoje, "duracao": "1:30", "descricao": "Reinício do AP", "faturavel": "1",
                       "tipo_servico_id": str(self.sql("SELECT id FROM tipos_servico WHERE nome = 'Help desk'")[0]["id"])})
        self.assertEqual(r.status_code, 302)
        self.post(eq, f"/chamados/{numero}/atualizar", {"status": "resolvido", "prioridade": "alta",
                                                       "contrato_id": str(self.ids["ct_beta"]),
                                                       "responsavel_id": str(self.ids["tec"])})

        # 5) admin aprova -> saldo do banco cai
        antes = self.sql("SELECT horas_saldo FROM vw_saldo_banco_horas WHERE contrato_id = %s",
                         [self.ids["ct_beta"]])[0]["horas_saldo"]
        ap = self.sql("SELECT id FROM apontamentos WHERE chamado_id = %s", [ch["id"]])[0]["id"]
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, "/horas/aprovar", {"ids": str(ap), "acao": "aprovado"})
        a = self.sql("SELECT status, aprovado_por FROM apontamentos WHERE id = %s", [ap])[0]
        self.assertEqual((a["status"], a["aprovado_por"]), ("aprovado", self.ids["adm"]))
        depois = self.sql("SELECT horas_saldo FROM vw_saldo_banco_horas WHERE contrato_id = %s",
                          [self.ids["ct_beta"]])[0]["horas_saldo"]
        self.assertEqual(antes, depois)  # pendente já consumia o saldo; aprovar não muda

        # 6) cliente avalia -> fechado
        self.post(cli, f"/portal/chamados/{numero}/avaliar", {"nota": "5", "comentario": "Rápido"})
        ch = self.sql("SELECT status, avaliacao FROM chamados WHERE numero = %s", [numero])[0]
        self.assertEqual((ch["status"], ch["avaliacao"]), ("fechado", 5))

        # auditoria registrou quem aprovou
        self.assertTrue(self.sql("""SELECT 1 FROM auditoria WHERE tabela = 'apontamentos' AND registro_id = %s
                                    AND acao = 'UPDATE' AND usuario_id = %s""", [ap, self.ids["adm"]]))

    def test_cadastro_cliente_e_contrato(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = self.post(adm, "/clientes/novo", {"razao_social": "Gama ME", "cnpj": "11.222.333/0001-82"})
        self.assertIn("CNPJ inválido".encode(), r.data)
        r = self.post(adm, "/clientes/novo", {"razao_social": "Gama ME", "cnpj": "", "dominio_email": "gama.com"})
        self.assertEqual(r.status_code, 302)
        cid = int(r.headers["Location"].rstrip("/").split("/")[-1])
        r = self.post(adm, f"/clientes/{cid}/contratos/novo",
                      {"tipo": "banco_horas", "descricao": "Pacote 10h", "data_inicio": "2026-10-01"})
        self.assertIn("horas do pacote".encode(), r.data)
        r = self.post(adm, f"/clientes/{cid}/contratos/novo",
                      {"tipo": "banco_horas", "descricao": "Pacote 10h", "data_inicio": "2026-10-01",
                       "horas_pacote": "10", "valor_hora": "150,00"})
        self.assertEqual(r.status_code, 302)
        ct = self.sql("SELECT * FROM contratos WHERE cliente_id = %s", [cid])[0]
        self.assertEqual(str(ct["valor_hora"]), "150.00")
        # contato do cliente recebe senha temporária e consegue entrar
        r = self.post(adm, f"/clientes/{cid}/contatos", {"nome": "Paula", "email": "paula@gama.com"},
                      follow_redirects=True)
        senha = re.search(rb"Senha tempor\xc3\xa1ria: (\S+) ", r.data).group(1).decode()
        c = self.cliente_http()
        r = self.post(c, "/entrar", {"email": "paula@gama.com", "senha": senha})
        self.assertEqual(r.headers["Location"], "/portal/")

    def test_horas_invalidas(self):
        eq = self.cliente_http("ana@portosolucoes.com.br")
        num = self.sql("SELECT numero FROM chamados LIMIT 1")[0]["numero"]
        r = self.post(eq, f"/chamados/{num}/horas", {"data": "2099-01-01", "duracao": "1h", "descricao": "x"},
                      follow_redirects=True)
        self.assertIn("futura".encode(), r.data)
        r = self.post(eq, f"/chamados/{num}/horas", {"data": "2026-10-01", "duracao": "abc", "descricao": "x"},
                      follow_redirects=True)
        self.assertIn("Informe a dura".encode(), r.data)

    def test_trocar_senha(self):
        c = self.cliente_http("joao@clinicabeta.com.br")
        r = self.post(c, "/conta/senha", {"atual": SENHA_DEMO, "nova": "curta", "confirmacao": "curta"})
        self.assertIn("10 caracteres".encode(), r.data)


if __name__ == "__main__":
    unittest.main()


class TestEmail(Base):
    TOKEN = "tok-teste"

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.app.config["INTEGRACAO_TOKEN"] = cls.TOKEN

    def chamar(self, rota, dados=None, token=None):
        return self.app.test_client().post(f"/integracao/email/{rota}", json=dados or {},
                                           headers={"X-Token": token or self.TOKEN})

    def test_token_obrigatorio(self):
        self.assertEqual(self.chamar("saida", token="errado").status_code, 403)

    def test_limpar_resposta(self):
        from app.email_integracao import limpar_resposta
        t = "Ok, funcionou!\n\nEm qui., 23 de out. de 2026 às 10:00, Suporte <suporte@x.com> escreveu:\n> antigo"
        self.assertEqual(limpar_resposta(t), "Ok, funcionou!")

    def test_fluxo_email(self):
        # cliente conhecido pelo domínio abre chamado por e-mail
        r = self.chamar("entrada", {"message_id": "m1", "thread_id": "t1", "de": "Carla@AlfaLog.com.br",
                                    "assunto": "Sem acesso à VPN", "texto": "Não conecto desde ontem."})
        self.assertEqual(r.json["acao"], "novo")
        numero = r.json["numero"]
        ch = self.sql("SELECT * FROM chamados WHERE numero = %s", [numero])[0]
        self.assertEqual((ch["origem"], ch["solicitante_email"], ch["cliente_id"]),
                         ("email", "carla@alfalog.com.br", self.ids["alfa"]))
        # duplicado é ignorado
        self.assertEqual(self.chamar("entrada", {"message_id": "m1", "de": "carla@alfalog.com.br"}).json["acao"],
                         "duplicado")
        # confirmação foi enfileirada e sai pela fila
        fila = self.chamar("saida").json["emails"]
        conf = [e for e in fila if e["para"] == "carla@alfalog.com.br"]
        self.assertTrue(conf and f"[#{numero}]" in conf[0]["assunto"])
        self.assertEqual(self.chamar("enviados", {"ids": [e["id"] for e in fila]}).json["confirmados"], len(fila))
        self.assertEqual(self.chamar("saida").json["emails"], [])
        # técnico responde no sistema -> notificação por e-mail
        eq = self.cliente_http("ana@portosolucoes.com.br")
        self.post(eq, f"/chamados/{numero}/responder", {"texto": "Reinicie o cliente VPN, por favor.",
                                                        "status_apos": "aguardando_cliente"})
        fila = self.chamar("saida").json["emails"]
        self.assertEqual(len(fila), 1)
        self.assertIn("Reinicie o cliente VPN", fila[0]["corpo"])
        self.chamar("enviados", {"ids": [fila[0]["id"]]})
        # cliente responde por e-mail (assunto com [#n]) -> vira mensagem e retoma atendimento
        r = self.chamar("entrada", {"message_id": "m2", "thread_id": "t1", "de": "carla@alfalog.com.br",
                                    "assunto": f"Re: [#{numero}] Sem acesso à VPN",
                                    "texto": "Funcionou!\n\nOn Thu, Oct 23, 2026 Suporte wrote:\n> Reinicie"})
        self.assertEqual(r.json["acao"], "resposta")
        msgs = self.sql("SELECT texto FROM chamado_interacoes c JOIN chamados x ON x.id = c.chamado_id "
                        "WHERE x.numero = %s ORDER BY c.id", [numero])
        self.assertEqual(msgs[-1]["texto"], "Funcionou!")
        self.assertEqual(self.sql("SELECT status FROM chamados WHERE numero = %s", [numero])[0]["status"],
                         "em_atendimento")
        # nota interna não gera e-mail
        self.post(eq, f"/chamados/{numero}/responder", {"texto": "nota", "interna": "1"})
        self.assertEqual(self.chamar("saida").json["emails"], [])

    def test_remetente_desconhecido(self):
        r = self.chamar("entrada", {"message_id": "x9", "de": "spam@qualquer.com", "assunto": "Oferta",
                                    "texto": "compre"})
        self.assertEqual(r.json["acao"], "recusado")
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertIn(b"spam@qualquer.com", adm.get("/emails-recusados").data)

    def test_intruso_nao_responde_chamado_alheio(self):
        num = self.sql("SELECT numero FROM chamados WHERE cliente_id = %s LIMIT 1", [self.ids["beta"]])[0]["numero"]
        r = self.chamar("entrada", {"message_id": "x10", "de": "maria@alfalog.com.br",
                                    "assunto": f"Re: [#{num}]", "texto": "oi"})
        self.assertEqual(r.json["acao"], "recusado")


class TestImportacao(Base):
    def planilha(self, clientes=(), contratos=(), contatos=()):
        import io
        from openpyxl import load_workbook
        from app.importacao import gerar_modelo
        wb = load_workbook(io.BytesIO(gerar_modelo()))
        for aba, linhas in (("Clientes", clientes), ("Contratos", contratos), ("Contatos", contatos)):
            ws = wb[aba]
            for col in range(1, ws.max_column + 1):
                ws.cell(row=2, column=col).value = None
            for i, l in enumerate(linhas, start=2):
                for j, v in enumerate(l, start=1):
                    ws.cell(row=i, column=j, value=v)
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return buf

    def enviar(self, arquivo):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        return self.post(adm, "/importar", {"planilha": (arquivo, "clientes.xlsx")},
                         content_type="multipart/form-data")

    def test_modelo(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = adm.get("/importar/modelo.xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data[:2] == b"PK")
        self.assertEqual(adm.get("/importar").status_code, 200)

    def test_importacao_ok(self):
        r = self.enviar(self.planilha(
            clientes=[("Delta Engenharia Ltda", "Delta", "33.445.566/0001-86", None, None, "delta.eng.br", None)],
            contratos=[("Delta", "Banco de horas", "Pacote 20h", None, None, "150,00", 20, None, "01/10/2026"),
                       ("Alfa Logística", "Hora avulsa", "Avulso", None, None, 200, None, None, "2026-10-01")],
            contatos=[("33445566000186", "Pedro", "pedro@delta.eng.br", None)]))
        self.assertIn("Importação concluída".encode(), r.data, r.data[-1500:])
        cid = self.sql("SELECT id FROM clientes WHERE cnpj = '33445566000186'")[0]["id"]
        self.assertEqual(self.sql("SELECT count(*) AS n FROM contratos WHERE cliente_id = %s", [cid])[0]["n"], 1)
        self.assertEqual(self.sql("SELECT horas_pacote FROM contratos WHERE cliente_id = %s", [cid])[0]["horas_pacote"], 20)
        self.assertTrue(self.sql("SELECT 1 FROM contratos WHERE cliente_id = %s AND tipo = 'hora_avulsa'",
                                 [self.ids["alfa"]]))
        self.assertTrue(self.sql("SELECT 1 FROM usuarios WHERE email = 'pedro@delta.eng.br' AND cliente_id = %s", [cid]))

    def test_importacao_com_erro_nao_grava(self):
        r = self.enviar(self.planilha(
            clientes=[("Ômega SA", "Ômega", None, None, None, None, None)],
            contratos=[("Ômega", "Mensal fixo", "Suporte", None, None, None, None, None, "01/10/2026")],
            contatos=[("Inexistente", "Zé", "ze@x.com", None)]))
        self.assertIn("Nada foi gravado".encode(), r.data)
        self.assertIn("preencha Valor mensal".encode(), r.data)
        self.assertIn("Inexistente".encode(), r.data)
        self.assertFalse(self.sql("SELECT 1 FROM clientes WHERE razao_social = 'Ômega SA'"))


class TestMigracoes(Base):
    def test_atualizar_idempotente(self):
        with self.app.app_context():
            from app.cli import atualizar_banco
            self.assertEqual(atualizar_banco(), [])
        self.assertTrue(self.sql("SELECT 1 FROM schema_migracoes WHERE arquivo = '020_email.sql'"))


class TestPoliticasSLA(Base):
    def test_politica_por_contrato(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        prazos = {"critica_resposta": "30m", "critica_solucao": "2h", "alta_resposta": "1h", "alta_solucao": "4h",
                  "media_resposta": "2h", "media_solucao": "10h", "baixa_resposta": "4h", "baixa_solucao": "20h"}
        r = self.post(adm, "/config/sla", {"acao": "criar", "nome": "Premium", **prazos}, follow_redirects=True)
        self.assertIn("Premium".encode(), r.data)
        pid = self.sql("SELECT id FROM sla_politicas WHERE nome = 'Premium'")[0]["id"]
        self.assertEqual(self.sql("SELECT minutos_solucao FROM sla_regras WHERE politica_id = %s AND prioridade = 'critica'",
                                  [pid])[0]["minutos_solucao"], 120)
        # erro: solução menor que resposta
        r = self.post(adm, "/config/sla", {"acao": "criar", "nome": "Ruim", **{**prazos, "alta_solucao": "30m"}},
                      follow_redirects=True)
        self.assertIn("não pode ser menor".encode(), r.data)
        # contrato da Beta passa a usar Premium
        ct = self.sql("SELECT * FROM contratos WHERE id = %s", [self.ids["ct_beta"]])[0]
        dados = {k: ("" if ct[k] is None else str(ct[k])) for k in
                 ["tipo", "descricao", "valor_mensal", "franquia_horas", "valor_hora", "horas_pacote", "valor_total",
                  "data_inicio", "data_fim", "dia_vencimento", "indice_reajuste", "mes_reajuste"]}
        self.post(adm, f"/contratos/{self.ids['ct_beta']}", {**dados, "ativo": "1", "sla_politica_id": str(pid)})
        self.assertEqual(self.sql("SELECT sla_politica_id FROM contratos WHERE id = %s", [self.ids["ct_beta"]])[0]
                         ["sla_politica_id"], pid)
        # política em uso não pode ser excluída
        r = self.post(adm, "/config/sla", {"acao": "excluir", "politica_id": str(pid)}, follow_redirects=True)
        self.assertIn("em uso".encode(), r.data)
        # chamado crítico da Beta usa 2h úteis (Premium) e não 4h (padrão)
        cli = self.cliente_http("joao@clinicabeta.com.br")
        r = self.post(cli, "/portal/chamados/novo", {"titulo": "Servidor fora", "descricao": "x", "prioridade": "critica"})
        num = int(r.headers["Location"].rstrip("/").split("/")[-1])
        c = self.sql("""SELECT fn_minutos_uteis_entre(criado_em, prazo_solucao) AS m,
                               fn_somar_minutos_uteis(criado_em, 120) = prazo_solucao AS ok
                          FROM chamados WHERE numero = %s""", [num])[0]
        self.assertTrue(c["ok"])


class TestLogo(Base):
    def test_logo_padrao_e_troca(self):
        import io
        c = self.app.test_client()
        r = c.get("/marca/logo")
        self.assertEqual((r.status_code, r.mimetype), (200, "image/png"))
        padrao = r.data
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertEqual(adm.get("/config/marca").status_code, 200)
        r = self.post(adm, "/config/marca", {"logo": (io.BytesIO(b"<svg>x</svg>"), "x.svg")},
                      content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Formato não aceito".encode(), r.data)
        novo = b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 4
        self.post(adm, "/config/marca", {"logo": (io.BytesIO(novo), "logo.png")}, content_type="multipart/form-data")
        self.assertEqual(c.get("/marca/logo").data, novo)
        self.post(adm, "/config/marca", {"acao": "restaurar"})
        self.assertEqual(c.get("/marca/logo").data, padrao)
        # técnico não troca logo
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(tec.get("/config/marca").status_code, 403)


class TestCriarAdmin(Base):
    def test_criar_admin(self):
        r = self.app.test_cli_runner().invoke(args=["criar-admin", "Novo.Admin@portosolucoes.com.br", "--nome", "Novo"])
        self.assertEqual(r.exit_code, 0, r.output)
        senha = r.output.strip().split(": ")[-1]
        self.assertEqual(len(senha), 12)
        c = self.cliente_http()
        r = self.post(c, "/entrar", {"email": "novo.admin@portosolucoes.com.br", "senha": f" {senha} "})
        self.assertEqual(r.headers["Location"], "/painel")
        # e-mail que já é usuário de cliente: cria um RECURSO admin separado (código R...), cliente intacto
        r = self.app.test_cli_runner().invoke(args=["criar-admin", "maria@alfalog.com.br"])
        contas = self.sql("SELECT perfil, codigo FROM usuarios WHERE email = 'maria@alfalog.com.br' ORDER BY perfil")
        self.assertEqual([c["perfil"] for c in contas], ["admin", "cliente"])
        self.assertTrue(contas[0]["codigo"].startswith("R") and contas[1]["codigo"].startswith("C"))
        # reaplicar no mesmo recurso só troca a senha
        self.app.test_cli_runner().invoke(args=["criar-admin", "maria@alfalog.com.br"])
        self.assertEqual(len(self.sql("SELECT 1 FROM usuarios WHERE email = 'maria@alfalog.com.br'")), 2)


class TestRecuperarSenha(Base):
    def test_fluxo_recuperacao(self):
        c = self.cliente_http()
        # botão "Esqueci minha senha" envia o e-mail digitado no login
        r = self.post(c, "/recuperar-senha", {"email": "Joao@ClinicaBeta.com.br", "senha": ""})
        self.assertEqual(r.status_code, 302)
        email = self.sql("SELECT * FROM emails_saida WHERE evento = 'senha' ORDER BY id DESC LIMIT 1")[0]
        self.assertEqual(email["para"], "joao@clinicabeta.com.br")
        link = re.search(r"(/redefinir-senha/\S+)", email["corpo"]).group(1)
        self.assertEqual(c.get(link).status_code, 200)
        r = self.post(c, link, {"nova": "curta", "confirmacao": "curta"})
        self.assertIn("10 caracteres".encode(), r.data)
        r = self.post(c, link, {"nova": "NovaSenha-2026!", "confirmacao": "NovaSenha-2026!"})
        self.assertEqual(r.status_code, 302)
        c.get("/entrar")
        r = self.post(c, "/entrar", {"email": "joao@clinicabeta.com.br", "senha": "NovaSenha-2026!"})
        self.assertEqual(r.headers["Location"], "/portal/")
        # link não pode ser reutilizado
        r = c.get(link)
        self.assertIn("/recuperar-senha", r.headers["Location"])

    def test_email_desconhecido_mesma_resposta(self):
        c = self.cliente_http()
        antes = self.sql("SELECT count(*) AS n FROM emails_saida")[0]["n"]
        r = self.post(c, "/recuperar-senha", {"email": "ninguem@x.com"}, follow_redirects=True)
        self.assertIn("Se o e-mail estiver cadastrado".encode(), r.data)
        self.assertEqual(self.sql("SELECT count(*) AS n FROM emails_saida")[0]["n"], antes)

    def test_limite_e_expiracao(self):
        c = self.cliente_http()
        for _ in range(5):
            self.post(c, "/recuperar-senha", {"email": "ana@portosolucoes.com.br"})
        n = self.sql("""SELECT count(*) AS n FROM senha_tokens t JOIN usuarios u ON u.id = t.usuario_id
                        WHERE u.email = 'ana@portosolucoes.com.br'""")[0]["n"]
        self.assertEqual(n, 3)
        self.sql("UPDATE senha_tokens SET expira_em = now() - interval '1 minute'")
        corpo = self.sql("SELECT corpo FROM emails_saida WHERE para = 'ana@portosolucoes.com.br' ORDER BY id DESC LIMIT 1")[0]["corpo"]
        link = re.search(r"(/redefinir-senha/\S+)", corpo).group(1)
        self.assertIn("/recuperar-senha", c.get(link).headers["Location"])

    def test_login_tem_botao(self):
        r = self.app.test_client().get("/entrar")
        self.assertIn("Esqueci minha senha".encode(), r.data)


class TestEntrega1(Base):
    def tipo(self, nome="Help desk"):
        return str(self.sql("SELECT id FROM tipos_servico WHERE nome = %s", [nome])[0]["id"])

    def test_codigo_e_mesmo_email(self):
        u = self.sql("SELECT codigo FROM usuarios WHERE email = 'contato@portosolucoes.com.br'")[0]
        self.assertRegex(u["codigo"], r"^R\d{4}$")
        adm = self.cliente_http("contato@portosolucoes.com.br")
        # contato@ também como usuário do cliente Alfa: permitido (mesmo e-mail, outro código)
        r = self.post(adm, f"/clientes/{self.ids['alfa']}/contatos",
                      {"nome": "Yuri (cliente)", "email": "contato@portosolucoes.com.br"}, follow_redirects=True)
        cod = re.search(rb"Usu\xc3\xa1rio (C\d{4})", r.data).group(1).decode()
        senha = re.search(rb"Senha tempor\xc3\xa1ria: (\S+) ", r.data).group(1).decode()
        # login por código entra no portal; repetido no mesmo cliente é recusado
        c = self.cliente_http()
        r = self.post(c, "/entrar", {"email": cod.lower(), "senha": senha})
        self.assertEqual(r.headers["Location"], "/portal/")
        r = self.post(adm, f"/clientes/{self.ids['alfa']}/contatos",
                      {"nome": "X", "email": "contato@portosolucoes.com.br"}, follow_redirects=True)
        self.assertIn("já é usuário deste cliente".encode(), r.data)
        # login pelo e-mail com a senha do recurso continua indo para o painel
        c2 = self.cliente_http()
        r = self.post(c2, "/entrar", {"email": "contato@portosolucoes.com.br", "senha": SENHA_DEMO})
        self.assertEqual(r.headers["Location"], "/painel")

    def test_dominio(self):
        from app.util import normalizar_dominio as n
        self.assertEqual(n("Contato@Empresa.COM.br"), "empresa.com.br")
        self.assertEqual(n("https://www.empresa.com.br/x"), "empresa.com.br")
        self.assertIsNone(n(""))
        for ruim in ("gmail.com", "empresa", "a@b"):
            with self.assertRaises(ValueError):
                n(ruim)
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = self.post(adm, "/clientes/novo", {"razao_social": "Zeta", "dominio_email": "@hotmail.com"})
        self.assertIn("provedor público".encode(), r.data)

    def test_usuario_cliente_inativo(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        cid = self.sql("INSERT INTO clientes (razao_social, ativo) VALUES ('Inativa', false) RETURNING id")[0]["id"]
        r = self.post(adm, f"/clientes/{cid}/contatos", {"nome": "A", "email": "a@inativa.com"}, follow_redirects=True)
        self.assertIn("cliente está inativo".encode(), r.data)
        r = adm.get("/usuarios-clientes")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn(b">Inativa<", r.data)  # não aparece na lista de clientes para vincular

    def test_acessos_por_rotina(self):
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(tec.get("/relatorios/apontamentos").status_code, 403)
        self.assertNotIn(b"Relat\xc3\xb3rio de apontamentos", tec.get("/painel").data)
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertEqual(adm.get("/equipe/acessos").status_code, 200)
        ana = self.ids["tec"]
        dados = {f"p_{ana}_{r}": "1" for r in ("chamados", "chamados_gerir", "horas", "emails", "relatorios", "conhecimento")}  # libera relatórios
        self.post(adm, "/equipe/acessos", dados)
        self.assertEqual(tec.get("/relatorios/apontamentos").status_code, 200)
        dados.pop(f"p_{ana}_chamados")  # bloqueia chamados
        self.post(adm, "/equipe/acessos", dados)
        self.assertEqual(tec.get("/chamados").status_code, 403)
        self.assertEqual(self.sql("SELECT count(*) AS n FROM permissoes WHERE usuario_id = %s", [ana])[0]["n"], 2)

    def test_reset_senha_pelo_admin(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = self.post(adm, f"/usuarios/{self.ids['tec']}", {"acao": "link"}, follow_redirects=True)
        self.assertIn("Link para criar nova senha enviado".encode(), r.data)
        self.assertTrue(self.sql("SELECT 1 FROM emails_saida WHERE para = 'ana@portosolucoes.com.br' AND evento = 'senha'"))
        r = self.post(adm, f"/usuarios/{self.ids['maria']}", {"acao": "senha"}, follow_redirects=True)
        self.assertIn("Nova senha temporária".encode(), r.data)
        # técnico não reseta senha de ninguém
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(self.post(tec, f"/usuarios/{self.ids['adm']}", {"acao": "senha"}).status_code, 403)

    def test_tipos_horas_contrato_relatorio(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, "/config/tipos-servico", {"acao": "criar", "nome": "Desenvolvimento web"})
        dev = self.tipo("Desenvolvimento web")
        tec = self.cliente_http("ana@portosolucoes.com.br")
        import datetime as dt
        hoje = dt.datetime.now(util.TZ).date().isoformat()
        r = self.post(tec, "/horas/lancar", {"contrato_id": str(self.ids["ct_beta"]), "tipo_servico_id": dev,
                                             "data": hoje, "duracao": "3h", "descricao": "Integração API",
                                             "faturavel": "1"}, follow_redirects=True)
        self.assertIn("Horas lançadas no contrato".encode(), r.data)
        a = self.sql("SELECT * FROM apontamentos WHERE descricao = 'Integração API'")[0]
        self.assertIsNone(a["chamado_id"])
        # saldo de horas liberadas desconta
        saldo = self.sql("SELECT horas_consumidas FROM vw_saldo_banco_horas WHERE contrato_id = %s",
                         [self.ids["ct_beta"]])[0]["horas_consumidas"]
        self.assertGreaterEqual(saldo, 3)
        # painel do cliente mostra horas por tipo
        cli = self.cliente_http("joao@clinicabeta.com.br")
        self.assertIn("Desenvolvimento web".encode(), cli.get("/portal/").data)
        # relatório + Excel
        r = adm.get(f"/relatorios/apontamentos?cliente={self.ids['beta']}&tipo={dev}")
        self.assertIn("Integração API".encode(), r.data)
        x = adm.get(f"/relatorios/apontamentos?cliente={self.ids['beta']}&formato=xlsx")
        self.assertEqual(x.data[:2], b"PK")
        # tipo obrigatório no lançamento
        r = self.post(tec, "/horas/lancar", {"contrato_id": str(self.ids["ct_beta"]), "data": hoje, "duracao": "1h",
                                             "descricao": "x"}, follow_redirects=True)
        self.assertIn("tipo de serviço".encode(), r.data)

    def test_anexos(self):
        import io, zipfile
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("log.txt", "erro 42")
        zipdados = buf.getvalue()
        cli = self.cliente_http("joao@clinicabeta.com.br")
        r = self.post(cli, "/portal/chamados/novo", {"titulo": "Com anexo", "descricao": "veja", "prioridade": "baixa",
                                                     "anexo": (io.BytesIO(zipdados), "logs.zip")},
                      content_type="multipart/form-data")
        num = int(r.headers["Location"].rstrip("/").split("/")[-1])
        an = self.sql("""SELECT a.id, a.tamanho_bytes FROM chamado_anexos a JOIN chamados c ON c.id = a.chamado_id
                         WHERE c.numero = %s""", [num])[0]
        self.assertEqual(an["tamanho_bytes"], len(zipdados))
        self.assertEqual(cli.get(f"/anexos/{an['id']}").data, zipdados)
        # outro cliente não baixa
        outro = self.cliente_http("maria@alfalog.com.br")
        self.assertEqual(outro.get(f"/anexos/{an['id']}").status_code, 404)
        # tipo não compactado é recusado (e o chamado não é criado)
        r = self.post(cli, "/portal/chamados/novo", {"titulo": "PDF", "descricao": "x", "prioridade": "baixa",
                                                     "anexo": (io.BytesIO(b"%PDF-1.4"), "a.pdf")},
                      content_type="multipart/form-data")
        self.assertIn("compactados".encode(), r.data)
        self.assertFalse(self.sql("SELECT 1 FROM chamados WHERE titulo = 'PDF'"))
        # extensão .zip falsa também é recusada
        r = self.post(cli, f"/portal/chamados/{num}/responder", {"texto": "x",
                      "anexo": (io.BytesIO(b"nao sou zip"), "falso.zip")}, content_type="multipart/form-data",
                      follow_redirects=True)
        self.assertIn("não parece ser um arquivo compactado".encode(), r.data)
        # acima de 20 MB: recusado com orientação do Drive
        grande = b"PK\x03\x04" + b"0" * (20 * 1024 * 1024 + 10)
        r = self.post(cli, f"/portal/chamados/{num}/responder", {"texto": "y",
                      "anexo": (io.BytesIO(grande), "grande.zip")}, content_type="multipart/form-data",
                      follow_redirects=True)
        self.assertIn("Google Drive".encode(), r.data)
        self.assertEqual(self.sql("SELECT count(*) AS n FROM chamado_anexos a JOIN chamados c ON c.id = a.chamado_id "
                                  "WHERE c.numero = %s", [num])[0]["n"], 1)

    def test_menu_celular(self):
        cli = self.cliente_http("joao@clinicabeta.com.br")
        html = cli.get("/portal/").data
        self.assertIn(b'class="menu-btn"', html)
        self.assertIn(b'id="menu"', html)


class TestEntrega2(Base):
    TOKEN = "tok2"

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.app.config["INTEGRACAO_TOKEN"] = cls.TOKEN

    def chamar(self, rota, dados=None):
        return self.app.test_client().post(rota, json=dados or {}, headers={"X-Token": self.TOKEN})

    def test_relatorios(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.sql("UPDATE usuarios SET custo_hora = 80 WHERE id = %s", [self.ids["tec"]])
        r = adm.get("/relatorios/")
        self.assertEqual(r.status_code, 200)
        for chave in ("sla", "chamados", "produtividade", "contratos", "satisfacao", "prefaturamento", "rentabilidade"):
            with self.subTest(chave=chave):
                r = adm.get(f"/relatorios/{chave}?de=2026-01-01")
                self.assertEqual(r.status_code, 200, r.data[-800:])
                x = adm.get(f"/relatorios/{chave}?de=2026-01-01&formato=xlsx")
                self.assertEqual(x.data[:2], b"PK")
        # pré-faturamento: mensalidade Alfa 2.500 no mês corrente
        r = adm.get(f"/relatorios/prefaturamento?cliente={self.ids['alfa']}")
        self.assertIn("2.500,00".encode(), r.data)
        # técnico sem relatórios financeiros
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(tec.get("/relatorios/rentabilidade").status_code, 403)

    def test_cadastros(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = self.post(adm, "/clientes/novo", {"razao_social": "Ômega Ltda", "cep": "01310-100", "uf": "SP",
                                              "cidade": "São Paulo", "segmento": "Saúde",
                                              "responsavel_id": str(self.ids["tec"])})
        self.assertEqual(r.status_code, 302)
        c = self.sql("SELECT * FROM clientes WHERE razao_social = 'Ômega Ltda'")[0]
        self.assertEqual((c["cep"], c["uf"], c["responsavel_id"]), ("01310100", "SP", self.ids["tec"]))
        r = self.post(adm, f"/usuarios/{self.ids['tec']}/editar", {"nome": "Ana Técnica", "email": "ana@portosolucoes.com.br",
                                                                   "custo_hora": "92,50", "cargo": "Analista N2"})
        self.assertEqual(self.sql("SELECT custo_hora FROM usuarios WHERE id = %s", [self.ids["tec"]])[0]["custo_hora"],
                         __import__("decimal").Decimal("92.50"))
        self.post(adm, f"/usuarios/{self.ids['joao']}/editar", {"nome": "João Lima", "email": "joao@clinicabeta.com.br",
                                                                "recebe_nps": "1"})
        self.assertTrue(self.sql("SELECT recebe_nps FROM usuarios WHERE id = %s", [self.ids["joao"]])[0]["recebe_nps"])

    def test_nps(self):
        self.sql("UPDATE usuarios SET recebe_nps = true WHERE id = %s", [self.ids["joao"]])
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, "/config/integracoes", {"acao": "nps", "nps_ativo": "1", "nps_dia": "1"})
        r = self.chamar("/integracao/tarefas")
        self.assertGreaterEqual(r.json["nps_enviados"], 1)
        self.assertEqual(self.chamar("/integracao/tarefas").json["nps_enviados"], 0)  # 1x por mês
        corpo = self.sql("SELECT corpo FROM emails_saida WHERE evento = 'nps' ORDER BY id DESC LIMIT 1")[0]["corpo"]
        link = re.search(r"(/nps/\S+)", corpo).group(1)
        c = self.app.test_client()
        self.assertEqual(c.get(link).status_code, 200)
        r = self.post(c, link, {"nota": "9", "comentario": "Ótimo suporte"})
        self.assertIn("Obrigado".encode(), r.data)
        self.assertEqual(self.sql("SELECT nota FROM nps_respostas WHERE usuario_id = %s", [self.ids["joao"]])[0]["nota"], 9)
        r = adm.get("/relatorios/satisfacao?de=2026-01-01")
        self.assertIn("Ótimo suporte".encode(), r.data)

    def test_chat_alertas(self):
        from unittest import mock
        enviados = []
        class R:
            status_code = 200
        with mock.patch("app.integracoes.requests.post", side_effect=lambda url, json, timeout: enviados.append(json) or R()):
            adm = self.cliente_http("contato@portosolucoes.com.br")
            r = self.post(adm, "/config/integracoes", {"acao": "chat", "chat_webhook": "https://evil.com/x"},
                          follow_redirects=True)
            self.assertIn("chat.googleapis.com".encode(), r.data)
            self.post(adm, "/config/integracoes", {"acao": "chat",
                                                   "chat_webhook": "https://chat.googleapis.com/v1/spaces/X/messages?key=k"})
            cli = self.cliente_http("joao@clinicabeta.com.br")
            self.post(cli, "/portal/chamados/novo", {"titulo": "Tudo parado", "descricao": "x", "prioridade": "critica"})
            self.assertTrue(any("CRÍTICO" in e["text"] for e in enviados))
            # SLA estourado gera alerta 1 única vez (prazo forçado no passado: o teste não depende do dia/hora)
            self.sql("""UPDATE chamados SET prazo_solucao = now() - interval '2 hours', alerta_estouro_em = NULL,
                               alerta_risco_em = NULL
                         WHERE id = (SELECT min(id) FROM chamados WHERE status IN ('novo','em_atendimento'))""")
            n1 = self.chamar("/integracao/tarefas").json["alertas_sla"]
            n2 = self.chamar("/integracao/tarefas").json["alertas_sla"]
            self.assertGreaterEqual(n1, 1)
            self.assertEqual(n2, 0)

    def test_agenda_drive_bi(self):
        import io, zipfile, base64, datetime as dt
        adm = self.cliente_http("contato@portosolucoes.com.br")
        amanha = (dt.datetime.now(util.TZ) + dt.timedelta(days=1)).date().isoformat()
        num = self.sql("SELECT numero FROM chamados WHERE cliente_id = %s LIMIT 1", [self.ids["beta"]])[0]["numero"]
        r = self.post(adm, f"/chamados/{num}/agendar", {"tipo": "visita", "data": amanha, "hora": "10:00",
                                                        "duracao": "2h", "recurso_id": str(self.ids["tec"]),
                                                        "convidar_cliente": "1", "local": "Rua X, 10"},
                      follow_redirects=True)
        self.assertIn("Agendado".encode(), r.data)
        itens = self.chamar("/integracao/agenda").json["itens"]
        self.assertEqual(len(itens), 1)
        self.assertIn("joao@clinicabeta.com.br", itens[0]["convidados"])
        self.assertIn("ana@portosolucoes.com.br", itens[0]["convidados"])
        self.chamar("/integracao/agenda/ok", {"id": itens[0]["id"], "acao": "criado", "evento_id": "ev1"})
        self.assertEqual(self.chamar("/integracao/agenda").json["itens"], [])
        ag = self.sql("SELECT id FROM agendamentos")[0]["id"]
        self.post(adm, f"/agendamentos/{ag}/cancelar", {})
        self.assertEqual(self.chamar("/integracao/agenda").json["itens"][0]["status"], "cancelar")
        # drive
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("a.txt", "x")
        cli = self.cliente_http("joao@clinicabeta.com.br")
        self.post(cli, f"/portal/chamados/{num}/responder", {"texto": "segue", "anexo": (io.BytesIO(buf.getvalue()), "a.zip")},
                  content_type="multipart/form-data")
        a = self.chamar("/integracao/drive").json["anexo"]
        self.assertEqual(base64.b64decode(a["conteudo"]), buf.getvalue())
        self.chamar("/integracao/drive/ok", {"anexo_id": a["id"], "arquivo_id": "f1", "url": "https://drive/f1",
                                             "pasta_url": "https://drive/pasta"})
        self.assertIsNone(self.chamar("/integracao/drive").json["anexo"])
        self.assertIn(b"https://drive/pasta", adm.get(f"/chamados/{num}").data)
        # bi
        bi = self.chamar("/integracao/bi").json["abas"]
        self.assertIn("receita_horas", bi["apontamentos"]["colunas"])
        self.assertTrue(bi["chamados"]["linhas"])

    def test_2fa(self):
        from app.auth import totp_codigo
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(tec.get("/conta/2fa").status_code, 200)
        with tec.session_transaction() as s:
            segredo = s["totp_novo"]
        r = self.post(tec, "/conta/2fa", {"codigo": "000000"}, follow_redirects=True)
        self.assertIn("Código inválido".encode(), r.data)
        self.post(tec, "/conta/2fa", {"codigo": totp_codigo(segredo)})
        self.assertTrue(self.sql("SELECT totp_ativo FROM usuarios WHERE id = %s", [self.ids["tec"]])[0]["totp_ativo"])
        # login agora pede o código
        c = self.cliente_http()
        r = self.post(c, "/entrar", {"email": "ana@portosolucoes.com.br", "senha": SENHA_DEMO})
        self.assertIn("/entrar/2fa", r.headers["Location"])
        self.assertEqual(c.get("/painel").status_code, 302)  # ainda não entrou
        r = self.post(c, "/entrar/2fa", {"codigo": totp_codigo(segredo)})
        self.assertEqual(r.headers["Location"], "/painel")
        # admin desativa (celular perdido)
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, f"/usuarios/{self.ids['tec']}/editar", {"nome": "Ana Técnica", "email": "ana@portosolucoes.com.br",
                                                               "desativar_2fa": "1"})
        self.assertFalse(self.sql("SELECT totp_ativo FROM usuarios WHERE id = %s", [self.ids["tec"]])[0]["totp_ativo"])

    def test_telas_novas(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        num = self.sql("SELECT numero FROM chamados LIMIT 1")[0]["numero"]
        for url in ["/config/integracoes", f"/usuarios/{self.ids['tec']}/editar", f"/usuarios/{self.ids['maria']}/editar",
                    f"/chamados/{num}", "/clientes/novo", f"/clientes/{self.ids['alfa']}", "/equipe"]:
            with self.subTest(url=url):
                self.assertEqual(adm.get(url).status_code, 200)


class TestEntrega3(Base):
    """Cronômetro, saldos rateados, apontamentos pelo admin e administrador master."""

    def tipo(self):
        return str(self.sql("SELECT id FROM tipos_servico ORDER BY id LIMIT 1")[0]["id"])

    def chamado_beta(self, titulo):
        """Chamado no contrato 'Pacote 40h' (banco de horas) da Clínica Beta."""
        ct = self.sql("SELECT id FROM contratos WHERE descricao = 'Pacote 40h'")[0]["id"]
        r = self.sql("""INSERT INTO chamados (cliente_id, solicitante_id, contrato_id, titulo, descricao, prioridade, origem,
                                              responsavel_id)
                        VALUES (%s, %s, %s, %s, 'x', 'media', 'portal', %s) RETURNING id, numero""",
                     [self.ids["beta"], self.ids["joao"], ct, titulo, self.ids["tec"]])[0]
        return ct, r["id"], r["numero"]

    def test_cronometro(self):
        ct, cid, num = self.chamado_beta("Cronômetro")
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertIn("Iniciar", tec.get(f"/chamados/{num}").get_data(as_text=True))
        self.post(tec, f"/chamados/{num}/cronometro/iniciar", {})
        self.assertIsNotNone(self.sql("SELECT rodando_desde FROM cronometros WHERE chamado_id = %s", [cid])[0]["rodando_desde"])
        # simula 50 minutos corridos
        self.sql("UPDATE cronometros SET rodando_desde = now() - interval '50 minutes' WHERE chamado_id = %s", [cid])
        self.post(tec, f"/chamados/{num}/cronometro/pausar", {})
        c = self.sql("SELECT * FROM cronometros WHERE chamado_id = %s", [cid])[0]
        self.assertIsNone(c["rodando_desde"])
        self.assertGreaterEqual(c["segundos_acumulados"], 3000)
        self.assertIn("Retomar", tec.get(f"/chamados/{num}").get_data(as_text=True))
        # finalizar sem descrição não lança
        self.post(tec, f"/chamados/{num}/cronometro/finalizar", {"tipo_servico_id": self.tipo(), "descricao": ""})
        self.assertEqual(self.sql("SELECT count(*) AS n FROM apontamentos WHERE chamado_id = %s", [cid])[0]["n"], 0)
        self.post(tec, f"/chamados/{num}/cronometro/finalizar",
                  {"tipo_servico_id": self.tipo(), "descricao": "Ajuste no servidor", "faturavel": "1"})
        a = self.sql("SELECT * FROM apontamentos WHERE chamado_id = %s", [cid])
        self.assertEqual(len(a), 1)
        self.assertIn(a[0]["minutos"], (50, 51))
        self.assertEqual(a[0]["contrato_id"], ct)
        self.assertFalse(self.sql("SELECT 1 FROM cronometros WHERE chamado_id = %s", [cid]))

    def test_cronometro_respeita_saldo(self):
        ct, cid, num = self.chamado_beta("Limite")
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, "/admin/saldos", {"acao": "alocar", "chamado": str(num), "horas": "1", "contrato_id": str(ct)})
        self.assertEqual(float(self.sql("SELECT horas_alocadas FROM chamados WHERE id = %s", [cid])[0]["horas_alocadas"]), 1)
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.post(tec, f"/chamados/{num}/cronometro/iniciar", {})
        self.sql("UPDATE cronometros SET rodando_desde = now() - interval '3 hours' WHERE chamado_id = %s", [cid])
        self.post(tec, f"/chamados/{num}/cronometro/pausar", {})
        self.assertEqual(self.sql("SELECT segundos_acumulados FROM cronometros WHERE chamado_id = %s", [cid])[0]
                         ["segundos_acumulados"], 3600)
        html = tec.get(f"/chamados/{num}").get_data(as_text=True)
        self.assertIn("Saldo de horas esgotado", html)
        self.post(tec, f"/chamados/{num}/cronometro/iniciar", {})
        self.assertIsNone(self.sql("SELECT rodando_desde FROM cronometros WHERE chamado_id = %s", [cid])[0]["rodando_desde"])
        self.post(tec, f"/chamados/{num}/cronometro/finalizar", {"tipo_servico_id": self.tipo(), "descricao": "ok",
                                                                "faturavel": "1"})
        # lançamento manual além do saldo é bloqueado para não-admin
        r = self.post(tec, f"/chamados/{num}/horas", {"data": "2026-10-01", "duracao": "0:30", "descricao": "extra",
                                                       "tipo_servico_id": self.tipo(), "faturavel": "1"},
                      follow_redirects=True)
        self.assertEqual(self.sql("SELECT sum(minutos) AS s FROM apontamentos WHERE chamado_id = %s", [cid])[0]["s"], 60)
        self.assertIn("saldo", r.get_data(as_text=True).lower())

    def test_rateio_validado(self):
        ct, c1, n1 = self.chamado_beta("Rateio A")
        _, c2, n2 = self.chamado_beta("Rateio B")
        adm = self.cliente_http("contato@portosolucoes.com.br")
        livre = float(self.sql("SELECT horas_saldo FROM vw_saldo_banco_horas WHERE contrato_id = %s", [ct])[0]["horas_saldo"])
        reservado = float(self.sql("SELECT coalesce(sum(horas_alocadas),0) AS s FROM chamados WHERE contrato_id = %s",
                                   [ct])[0]["s"])
        self.post(adm, "/admin/saldos", {"acao": "alocar", "chamado": str(n1), "horas": "5", "contrato_id": str(ct)})
        r = self.post(adm, "/admin/saldos", {"acao": "alocar", "chamado": str(n2),
                                             "horas": str(int(livre - reservado)), "contrato_id": str(ct)},
                      follow_redirects=True)
        self.assertIn("Disponível para alocar", r.get_data(as_text=True))
        self.assertIsNone(self.sql("SELECT horas_alocadas FROM chamados WHERE id = %s", [c2])[0]["horas_alocadas"])
        # remover alocação
        self.post(adm, "/admin/saldos", {"acao": "alocar", "chamado": str(n1), "horas": "", "contrato_id": str(ct)})
        self.assertIsNone(self.sql("SELECT horas_alocadas FROM chamados WHERE id = %s", [c1])[0]["horas_alocadas"])
        # tipo do projeto e horas liberadas
        self.post(adm, "/admin/saldos", {"acao": "contrato", "contrato_id": str(ct), "horas_pacote": "60",
                                         "tipo_servico_id": self.tipo()})
        x = self.sql("SELECT horas_pacote, tipo_servico_id FROM contratos WHERE id = %s", [ct])[0]
        self.assertEqual(float(x["horas_pacote"]), 60)
        self.assertEqual(str(x["tipo_servico_id"]), self.tipo())
        self.assertEqual(adm.get(f"/admin/saldos?contrato={ct}").status_code, 200)

    def test_apontamentos_admin_e_proprio(self):
        _, cid, num = self.chamado_beta("Apontamentos")
        tec_id, adm_id = self.ids["tec"], self.ids["adm"]
        ap = self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                         VALUES (%s, %s, current_date, 30, 'original', %s) RETURNING id""",
                      [adm_id, cid, self.tipo()])[0]["id"]
        tec = self.cliente_http("ana@portosolucoes.com.br")
        # técnico não acessa telas de admin nem edita lançamento de outro
        self.assertEqual(tec.get("/admin/apontamentos").status_code, 403)
        self.assertEqual(tec.get("/admin/saldos").status_code, 403)
        self.assertEqual(tec.get(f"/horas/{ap}/editar").status_code, 403)
        self.assertEqual(self.post(tec, f"/admin/apontamentos/{ap}", {"descricao": "hack"}).status_code, 403)
        # técnico edita o próprio pendente
        meu = self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                          VALUES (%s, %s, current_date, 30, 'meu', %s) RETURNING id""", [tec_id, cid, self.tipo()])[0]["id"]
        self.post(tec, f"/horas/{meu}/editar", {"data": "2026-10-01", "duracao": "0:45", "descricao": "meu editado",
                                               "tipo_servico_id": self.tipo(), "faturavel": "1"})
        self.assertEqual(self.sql("SELECT minutos FROM apontamentos WHERE id = %s", [meu])[0]["minutos"], 45)
        # admin edita o de outro recurso, aprova e troca o recurso
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertEqual(adm.get("/admin/apontamentos").status_code, 200)
        self.assertEqual(adm.get(f"/admin/apontamentos/{meu}").status_code, 200)
        self.post(adm, f"/admin/apontamentos/{meu}", {"usuario_id": str(adm_id), "chamado": str(num), "status": "aprovado",
                                                     "data": "2026-10-02", "duracao": "1:00", "descricao": "corrigido",
                                                     "tipo_servico_id": self.tipo(), "faturavel": "1"})
        x = self.sql("SELECT * FROM apontamentos WHERE id = %s", [meu])[0]
        self.assertEqual((x["usuario_id"], x["minutos"], x["status"], x["aprovado_por"]), (adm_id, 60, "aprovado", adm_id))
        # técnico já não edita (aprovado)
        tec.get(f"/horas/{meu}/editar")
        self.assertEqual(self.sql("SELECT descricao FROM apontamentos WHERE id = %s", [meu])[0]["descricao"], "corrigido")
        # admin gera apontamento para o técnico em um contrato
        ct = self.sql("SELECT id FROM contratos WHERE descricao = 'Pacote 40h'")[0]["id"]
        self.assertEqual(adm.get("/admin/apontamentos/novo").status_code, 200)
        self.post(adm, "/admin/apontamentos/novo", {"usuario_id": str(tec_id), "chamado": "", "contrato_id": str(ct),
                                                   "status": "pendente", "data": "2026-10-03", "duracao": "2h",
                                                   "descricao": "gerado pelo admin", "tipo_servico_id": self.tipo()})
        g_ = self.sql("SELECT * FROM apontamentos WHERE descricao = 'gerado pelo admin'")[0]
        self.assertEqual((g_["usuario_id"], g_["minutos"], g_["contrato_id"], g_["faturavel"]), (tec_id, 120, ct, False))
        # admin exclui
        self.post(adm, f"/admin/apontamentos/{ap}", {"acao": "excluir"})
        self.assertFalse(self.sql("SELECT 1 FROM apontamentos WHERE id = %s", [ap]))

    def test_admin_master(self):
        adm_id = self.ids["adm"]
        outro = self.sql("""INSERT INTO usuarios (nome, email, perfil, senha_hash)
                            SELECT 'Outro Admin', 'outro@portosolucoes.com.br', 'admin', senha_hash FROM usuarios
                             WHERE id = %s RETURNING id""", [adm_id])[0]["id"]
        fin = self.sql("""INSERT INTO usuarios (nome, email, perfil, senha_hash)
                          SELECT 'Fin', 'fin@portosolucoes.com.br', 'financeiro', senha_hash FROM usuarios
                           WHERE id = %s RETURNING id""", [adm_id])[0]["id"]
        o = self.cliente_http("outro@portosolucoes.com.br")
        # outro admin não rebaixa, não desativa e não edita o master
        self.post(o, f"/usuarios/{adm_id}", {"acao": "perfil", "perfil": "tecnico"})
        self.post(o, f"/usuarios/{adm_id}", {"acao": "alternar"})
        self.assertEqual(o.get(f"/usuarios/{adm_id}/editar").status_code, 403)
        x = self.sql("SELECT perfil, ativo FROM usuarios WHERE id = %s", [adm_id])[0]
        self.assertEqual((x["perfil"], x["ativo"]), ("admin", True))
        # admin promove; não-admin não
        self.post(o, f"/usuarios/{fin}", {"acao": "perfil", "perfil": "admin"})
        self.assertEqual(self.sql("SELECT perfil FROM usuarios WHERE id = %s", [fin])[0]["perfil"], "admin")
        self.sql("UPDATE usuarios SET perfil = 'financeiro' WHERE id = %s", [fin])
        self.sql("""INSERT INTO permissoes (usuario_id, rotina, permitido) VALUES (%s, 'equipe', true)
                    ON CONFLICT DO NOTHING""", [fin])
        t = self.cliente_http("ana@portosolucoes.com.br")
        f = self.cliente_http("fin@portosolucoes.com.br")
        for c in (t, f):
            self.post(c, f"/usuarios/{self.ids['tec']}", {"acao": "perfil", "perfil": "admin"})
            self.post(c, f"/usuarios/{fin}", {"acao": "perfil", "perfil": "admin"})
        self.assertEqual(self.sql("SELECT perfil FROM usuarios WHERE id = %s", [self.ids["tec"]])[0]["perfil"], "tecnico")
        self.assertEqual(self.sql("SELECT perfil FROM usuarios WHERE id = %s", [fin])[0]["perfil"], "financeiro")
        self.assertNotIn('value="admin"', f.get("/equipe").get_data(as_text=True))  # não-admin nem vê a opção
        # master não troca o próprio e-mail
        m = self.cliente_http("contato@portosolucoes.com.br")
        self.post(m, f"/usuarios/{adm_id}/editar", {"nome": "Yuri", "email": "yuri@portosolucoes.com.br"})
        self.assertEqual(self.sql("SELECT email FROM usuarios WHERE id = %s", [adm_id])[0]["email"],
                         "contato@portosolucoes.com.br")
        # banco-atualizar restaura o master se alguém mexer direto no banco
        self.sql("UPDATE usuarios SET perfil = 'tecnico' WHERE id = %s", [adm_id])
        r = self.app.test_cli_runner().invoke(args=["banco-atualizar"])
        self.assertIn("restaurado", r.output)
        self.assertEqual(self.sql("SELECT perfil FROM usuarios WHERE id = %s", [adm_id])[0]["perfil"], "admin")


class TestEntrega4(Base):
    """Senha provisória, contratos, aprovação, privacidade, consultas SQL, dicionário, portal e menu."""

    def tipo(self):
        return str(self.sql("SELECT id FROM tipos_servico ORDER BY id LIMIT 1")[0]["id"])

    def novo_recurso(self, email, perfil="financeiro"):
        uid = self.sql("""INSERT INTO usuarios (nome, email, perfil, senha_hash)
                          SELECT %s, %s, %s, senha_hash FROM usuarios WHERE id = %s RETURNING id""",
                       [email.split("@")[0].title(), email, perfil, self.ids["adm"]])[0]["id"]
        return uid

    def test_senha_temporaria_cliente(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = self.post(adm, f"/usuarios/{self.ids['maria']}", {"acao": "senha"}, follow_redirects=True)
        temp = re.search(r"senha temporária de [^:]+: (\w+)", r.get_data(as_text=True), re.I).group(1)
        c = self.app.test_client()
        c.get("/entrar")
        r = self.post(c, "/entrar", {"email": "maria@alfalog.com.br", "senha": f" {temp} "})
        self.assertEqual(r.status_code, 302)
        r = c.get("/portal/", follow_redirects=False)
        self.assertIn("/conta/senha", r.headers.get("Location", ""))
        tela = c.get("/conta/senha").get_data(as_text=True)
        self.assertIn("Crie a sua senha", tela)
        self.assertNotIn('name="atual"', tela)
        r = self.post(c, "/conta/senha", {"nova": "MinhaSenhaNova123", "confirmacao": "MinhaSenhaNova123"})
        self.assertEqual(r.status_code, 302)
        self.assertFalse(self.sql("SELECT senha_provisoria FROM usuarios WHERE id = %s", [self.ids["maria"]])[0]
                         ["senha_provisoria"])
        self.assertEqual(c.get("/portal/").status_code, 200)
        # troca normal exige a senha atual (e aceita espaços colados)
        r = self.post(c, "/conta/senha", {"atual": "errada", "nova": "OutraSenha12345", "confirmacao": "OutraSenha12345"})
        self.assertIn("Senha atual incorreta", r.get_data(as_text=True))
        r = self.post(c, "/conta/senha", {"atual": " MinhaSenhaNova123 ", "nova": "OutraSenha12345",
                                          "confirmacao": "OutraSenha12345"})
        self.assertEqual(r.status_code, 302)
        c2 = self.app.test_client()
        c2.get("/entrar")
        self.assertEqual(self.post(c2, "/entrar", {"email": "maria@alfalog.com.br",
                                                   "senha": "OutraSenha12345"}).status_code, 302)

    def test_contratos(self):
        fin_id = self.novo_recurso("fin4@portosolucoes.com.br")
        fin = self.cliente_http("fin4@portosolucoes.com.br")
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(tec.get("/contratos").status_code, 403)
        self.assertEqual(fin.get("/contratos").status_code, 200)
        self.assertIn("Pacote 40h", fin.get("/contratos?situacao=todos").get_data(as_text=True))
        self.assertEqual(fin.get("/contratos/novo").status_code, 200)
        r = self.post(fin, "/contratos/novo", {"cliente_id": str(self.ids["alfa"]), "tipo": "hora_avulsa",
                                               "descricao": "Avulso 2026", "valor_hora": "200",
                                               "data_inicio": "2026-10-01"})
        self.assertEqual(r.status_code, 302)
        ct = self.sql("SELECT id FROM contratos WHERE descricao = 'Avulso 2026'")[0]["id"]
        volta = "/contratos?situacao=todos"
        r = self.post(fin, f"/contratos/{ct}", {"tipo": "hora_avulsa", "descricao": "Avulso 2026 b", "valor_hora": "210",
                                               "data_inicio": "2026-10-01", "ativo": "1", "voltar": volta})
        self.assertEqual(r.headers["Location"], volta)
        # bloqueando a rotina de contratos para o financeiro
        self.sql("INSERT INTO permissoes (usuario_id, rotina, permitido) VALUES (%s, 'contratos', false)", [fin_id])
        self.assertEqual(fin.get("/contratos").status_code, 403)
        self.assertEqual(fin.get(f"/contratos/{ct}").status_code, 403)

    def test_aprovacao_admin(self):
        cid = self.sql("SELECT id FROM chamados ORDER BY id LIMIT 1")[0]["id"]
        ids = [self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                           VALUES (%s, %s, current_date, 30, %s, %s) RETURNING id""",
                        [self.ids["tec"], cid, f"ap{i}", self.tipo()])[0]["id"] for i in range(2)]
        self.novo_recurso("fin5@portosolucoes.com.br")
        for email in ("ana@portosolucoes.com.br", "fin5@portosolucoes.com.br"):
            c = self.cliente_http(email)
            self.assertEqual(c.get("/admin/aprovacao").status_code, 403)
            self.assertEqual(self.post(c, "/horas/aprovar", {"ids": str(ids[0]), "acao": "aprovado"}).status_code, 403)
        adm = self.cliente_http("contato@portosolucoes.com.br")
        tela = adm.get("/admin/aprovacao").get_data(as_text=True)
        self.assertIn("ap0", tela)
        self.post(adm, "/admin/aprovacao", {"ids": str(ids[1]), "acao": "rejeitado"})
        self.assertEqual(self.sql("SELECT status FROM apontamentos WHERE id = %s", [ids[1]])[0]["status"], "pendente")
        self.post(adm, "/admin/aprovacao", {"ids": str(ids[1]), "acao": "rejeitado", "motivo": "Duplicado"})
        self.post(adm, "/admin/aprovacao", {"ids": str(ids[0]), "acao": "aprovado"})
        x = {r["id"]: r for r in self.sql("SELECT id, status, motivo_rejeicao FROM apontamentos WHERE id IN (%s,%s)", ids)}
        self.assertEqual((x[ids[0]]["status"], x[ids[1]]["status"], x[ids[1]]["motivo_rejeicao"]),
                         ("aprovado", "rejeitado", "Duplicado"))
        tec = self.cliente_http("ana@portosolucoes.com.br")
        hoje = self.sql("SELECT current_date AS d")[0]["d"]
        self.assertIn("Duplicado", tec.get(f"/horas?semana={hoje}").get_data(as_text=True))

    def test_privacidade(self):
        tec = self.cliente_http("ana@portosolucoes.com.br")
        for url in ["/equipe", "/equipe/acessos", f"/usuarios/{self.ids['adm']}/editar", "/admin/apontamentos",
                    "/admin/consultas", "/admin/banco", "/relatorios/produtividade"]:
            with self.subTest(url=url):
                self.assertEqual(tec.get(url).status_code, 403)
        cid = self.sql("SELECT id FROM chamados ORDER BY id LIMIT 1")[0]["id"]
        self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                    VALUES (%s, %s, current_date, 45, 'segredo do admin', %s)""", [self.ids["adm"], cid, self.tipo()])
        self.sql("INSERT INTO permissoes (usuario_id, rotina, permitido) VALUES (%s, 'relatorios', true)", [self.ids["tec"]])
        r = tec.get(f"/relatorios/apontamentos?recurso={self.ids['adm']}&status=todos").get_data(as_text=True)
        self.assertNotIn("segredo do admin", r)
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertIn("segredo do admin", adm.get("/relatorios/apontamentos?status=todos").get_data(as_text=True))
        # o chamado continua mostrando as horas de todos (informação do ticket)
        num = self.sql("SELECT numero FROM chamados WHERE id = %s", [cid])[0]["numero"]
        self.assertIn("segredo do admin", tec.get(f"/chamados/{num}").get_data(as_text=True))

    def test_consultas_sql(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = self.post(adm, "/admin/consultas", {"sql": "SELECT nome, email, senha_hash FROM usuarios WHERE email LIKE '%%porto%%' ORDER BY id;"})
        html = r.get_data(as_text=True)
        self.assertIn("contato@portosolucoes.com.br", html)
        self.assertIn("•••", html)
        self.assertNotIn("scrypt:", html)
        n = self.sql("SELECT count(*) AS n FROM clientes")[0]["n"]
        for sql in ["DELETE FROM clientes", "SELECT 1; DELETE FROM clientes",
                    "WITH x AS (DELETE FROM clientes RETURNING 1) SELECT count(*) FROM x",
                    "SELECT set_config('x.y', 'z', false); UPDATE clientes SET ativo = false"]:
            with self.subTest(sql=sql):
                html = self.post(adm, "/admin/consultas", {"sql": sql}).get_data(as_text=True)
                self.assertIn("aviso erro", html)
        self.assertEqual(self.sql("SELECT count(*) AS n FROM clientes")[0]["n"], n)
        r = self.post(adm, "/admin/consultas", {"sql": "SELECT numero, titulo FROM chamados", "acao": "excel"})
        self.assertEqual(r.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.post(adm, "/admin/consultas", {"sql": "SELECT 1 AS um", "acao": "salvar", "nome": "Teste"})
        sid = self.sql("SELECT id FROM consultas_salvas WHERE nome = 'Teste'")[0]["id"]
        self.assertIn("SELECT 1 AS um", adm.get(f"/admin/consultas?salva={sid}").get_data(as_text=True))
        self.assertGreaterEqual(self.sql("SELECT count(*) AS n FROM consultas_log")[0]["n"], 5)
        # a sessão segue funcionando depois das consultas
        self.assertEqual(adm.get("/painel").status_code, 200)

    def test_dicionario_e_pdfs(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        html = adm.get("/admin/banco").get_data(as_text=True)
        self.assertIn("apontamentos", html)
        self.assertIn("Horas trabalhadas pelos recursos", html)
        r = adm.get("/admin/banco?formato=pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        for url in ["/relatorios/sla?formato=pdf", "/relatorios/apontamentos?formato=pdf",
                    "/relatorios/prefaturamento?formato=pdf"]:
            with self.subTest(url=url):
                self.assertTrue(adm.get(url).data.startswith(b"%PDF"))
        runner = self.app.test_cli_runner()
        import tempfile, pathlib
        with tempfile.TemporaryDirectory() as d:
            arq = pathlib.Path(d) / "dic.pdf"
            res = runner.invoke(args=["dicionario-pdf", str(arq)])
            self.assertIn("gerado", res.output)
            self.assertTrue(arq.read_bytes().startswith(b"%PDF"))

    def test_portal_horas(self):
        cid = self.sql("SELECT id FROM chamados WHERE cliente_id = %s ORDER BY id LIMIT 1", [self.ids["beta"]])[0]["id"]
        self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                    VALUES (%s, %s, current_date, 75, 'Configuração do backup', %s)""", [self.ids["tec"], cid, self.tipo()])
        c = self.cliente_http("joao@clinicabeta.com.br")
        html = c.get("/portal/horas").get_data(as_text=True)
        self.assertIn("Configuração do backup", html)
        self.assertIn("Chamados e apontamentos", html)
        self.assertIn("Resumo por chamado", c.get("/portal/horas?visao=sintetica").get_data(as_text=True))
        self.assertTrue(c.get("/portal/horas?formato=pdf").data.startswith(b"%PDF"))
        r = c.get("/portal/horas?formato=xlsx")
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ["Por chamado", "Apontamentos", "Por tipo de serviço"])
        self.assertIn("Configuração do backup", [x.value for x in wb["Apontamentos"]["F"]])
        # cliente de outra empresa não vê
        m = self.cliente_http("maria@alfalog.com.br")
        self.assertNotIn("Configuração do backup", m.get("/portal/horas").get_data(as_text=True))

    def test_rotinas_no_cadastro_do_recurso(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        tec_id = self.ids["tec"]
        html = adm.get(f"/usuarios/{tec_id}/editar").get_data(as_text=True)
        self.assertIn('name="rotinas"', html)
        self.assertIn("só administrador", html)
        u = self.sql("SELECT nome, email FROM usuarios WHERE id = %s", [tec_id])[0]
        self.post(adm, f"/usuarios/{tec_id}/editar", {"nome": u["nome"], "email": u["email"], "perfil": "tecnico",
                                                     "rotinas": ["horas", "contratos", "emails"]})
        perm = {r["rotina"]: r["permitido"] for r in self.sql("SELECT rotina, permitido FROM permissoes WHERE usuario_id = %s",
                                                               [tec_id])}
        self.assertEqual(perm.get("contratos"), True)
        self.assertEqual(perm.get("chamados"), False)
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(tec.get("/contratos").status_code, 200)
        self.assertEqual(tec.get("/chamados").status_code, 403)
        # admin: todas marcadas e travadas
        outro = self.novo_recurso("adm4@portosolucoes.com.br", "admin")
        html = adm.get(f"/usuarios/{outro}/editar").get_data(as_text=True)
        self.assertEqual(html.count('name="rotinas"'), html.count('checked disabled'))
        self.assertIn("todas</b> as rotinas", html)
        # restaura
        self.sql("DELETE FROM permissoes WHERE usuario_id = %s", [tec_id])

    def test_menu_e_dashboard(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        html = adm.get("/painel").get_data(as_text=True)
        for pasta in ["Dashboards", "Atendimento", "Horas", "Relatórios", "Cadastros", "Configurações", "Administração"]:
            self.assertIn(f'data-pasta="{pasta}"', html)
        self.assertEqual(adm.get("/dashboard").status_code, 200)
        self.novo_recurso("fin6@portosolucoes.com.br")
        fin = self.cliente_http("fin6@portosolucoes.com.br")
        self.assertEqual(fin.get("/dashboard").status_code, 200)
        self.assertNotIn('data-pasta="Administração"', fin.get("/painel").get_data(as_text=True))
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertEqual(tec.get("/dashboard").status_code, 403)
        cli = self.cliente_http("joao@clinicabeta.com.br")
        self.assertIn('data-pasta="Relatórios"', cli.get("/portal/").get_data(as_text=True))


class TestEntrega5(Base):
    """Itens do contrato, validações, regras de atendimento, cronômetros concorrentes, e-mails de aprovação,
    filtros e exportações."""

    def tipo(self, nome):
        return self.sql("SELECT id FROM tipos_servico WHERE nome = %s", [nome])[0]["id"]

    def contrato(self, adm, **extra):
        base = {"cliente_id": str(self.ids["alfa"]), "tipo": "banco_horas", "descricao": "Itens", "valor_hora": "150",
                "data_inicio": "2026-01-01"}
        base.update(extra)
        return self.post(adm, "/contratos/novo", base, follow_redirects=True)

    def test_contrato_validacoes_e_itens(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        hd, dev = str(self.tipo("Help desk")), str(self.tipo("Desenvolvimento"))
        casos = [({"data_fim": "2025-12-01", "horas_pacote": "10"}, "anterior à data de início"),
                 ({"dia_vencimento": "32", "horas_pacote": "10"}, "de 1 a 31"),
                 ({"mes_reajuste": "13", "horas_pacote": "10"}, "de 1 a 12"),
                 ({"item_tipo": [hd, hd], "item_horas": ["5", "5"]}, "só pode aparecer uma vez"),
                 ({"horas_pacote": "8", "item_tipo": [hd, dev], "item_horas": ["5", "5"]}, "passa das horas liberadas")]
        for extra, msg in casos:
            with self.subTest(msg=msg):
                self.assertIn(msg, self.contrato(adm, descricao="Inválido", **extra).get_data(as_text=True))
        self.assertFalse(self.sql("SELECT 1 FROM contratos WHERE descricao = 'Inválido'"))
        # horas liberadas = soma dos itens
        self.contrato(adm, descricao="Com itens", item_tipo=[hd, dev], item_horas=["10", "30"],
                      item_valor=["", "220"], item_desc=["", "sprints"])
        ct = self.sql("SELECT id, horas_pacote, tipo_servico_id FROM contratos WHERE descricao = 'Com itens'")[0]
        self.assertEqual(float(ct["horas_pacote"]), 40)
        self.assertEqual(ct["tipo_servico_id"], int(hd))
        itens = self.sql("SELECT tipo_servico_id, horas, valor_hora FROM contrato_itens WHERE contrato_id = %s ORDER BY ordem",
                         [ct["id"]])
        self.assertEqual([(i["tipo_servico_id"], float(i["horas"])) for i in itens], [(int(hd), 10), (int(dev), 30)])
        self.assertEqual(float(itens[1]["valor_hora"]), 220)
        self.assertIn("sprints", adm.get(f"/contratos/{ct['id']}").get_data(as_text=True))
        # apontamento com tipo fora do contrato é recusado
        tec = self.cliente_http("ana@portosolucoes.com.br")
        hoje = self.sql("SELECT current_date AS d")[0]["d"].isoformat()
        r = self.post(tec, "/horas/lancar", {"contrato_id": str(ct["id"]), "data": hoje, "duracao": "1h",
                                             "descricao": "infra", "tipo_servico_id": str(self.tipo("Infraestrutura"))},
                      follow_redirects=True)
        self.assertIn("não previsto no contrato", r.get_data(as_text=True))
        self.post(tec, "/horas/lancar", {"contrato_id": str(ct["id"]), "data": hoje, "duracao": "1h",
                                         "descricao": "ok dev", "tipo_servico_id": dev, "faturavel": "1"})
        self.assertTrue(self.sql("SELECT 1 FROM apontamentos WHERE descricao = 'ok dev'"))

    def test_mensal_fixo_saldo_do_mes(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.contrato(adm, tipo="mensal_fixo", descricao="Mensal 10h", valor_mensal="1500", franquia_horas="10",
                      horas_pacote="10")
        ct = self.sql("SELECT id FROM contratos WHERE descricao = 'Mensal 10h'")[0]["id"]
        self.sql("""INSERT INTO apontamentos (usuario_id, contrato_id, data, minutos, descricao, tipo_servico_id)
                    VALUES (%s, %s, date_trunc('month', current_date)::date - 3, 540, 'mês passado', %s)""",
                 [self.ids["tec"], ct, self.tipo("Help desk")])
        s = self.sql("SELECT horas_saldo, periodo FROM vw_saldo_banco_horas WHERE contrato_id = %s", [ct])[0]
        self.assertEqual((float(s["horas_saldo"]), s["periodo"]), (10, "mensal"))

    def test_regras_de_atendimento(self):
        # cliente Alfa: dois contratos válidos -> o chamado nasce sem contrato
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.contrato(adm, descricao="Segundo Alfa", horas_pacote="5")
        maria = self.cliente_http("maria@alfalog.com.br")
        r = self.post(maria, "/portal/chamados/novo", {"titulo": "Sem contrato", "descricao": "x", "prioridade": "media"})
        num = int(re.search(r"/chamados/(\d+)", r.headers["Location"]).group(1))
        cid = self.sql("SELECT id, contrato_id FROM chamados WHERE numero = %s", [num])[0]
        self.assertIsNone(cid["contrato_id"])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        html = tec.get(f"/chamados/{num}").get_data(as_text=True)
        self.assertIn("Atendimento bloqueado", html)
        self.post(tec, f"/chamados/{num}/responder", {"texto": "olá"})
        self.post(tec, f"/chamados/{num}/cronometro/iniciar", {})
        self.assertFalse(self.sql("SELECT 1 FROM chamado_interacoes WHERE chamado_id = %s", [cid["id"]]))
        self.assertFalse(self.sql("SELECT 1 FROM cronometros WHERE chamado_id = %s", [cid["id"]]))
        # sem permissão de gerir, não altera contrato/responsável/status/prioridade
        self.sql("INSERT INTO permissoes (usuario_id, rotina, permitido) VALUES (%s, 'chamados_gerir', false)",
                 [self.ids["tec"]])
        outro = self.sql("SELECT id FROM contratos WHERE cliente_id = %s AND ativo ORDER BY id LIMIT 1", [self.ids["alfa"]])[0]["id"]
        self.post(tec, f"/chamados/{num}/atualizar", {"contrato_id": str(outro), "responsavel_id": str(self.ids["tec"]),
                                                     "status": "novo", "prioridade": "alta"})
        self.assertIsNone(self.sql("SELECT contrato_id FROM chamados WHERE id = %s", [cid["id"]])[0]["contrato_id"])
        html = tec.get(f"/chamados/{num}").get_data(as_text=True)
        self.assertIn('id="form-atualizar"', html)  # status, prioridade e responsável continuam editáveis
        self.assertIn('name="contrato_id" required disabled', html)
        self.sql("DELETE FROM permissoes WHERE usuario_id = %s AND rotina = 'chamados_gerir'", [self.ids["tec"]])
        # contrato sem saldo não é aceito
        sem = self.sql("""INSERT INTO contratos (cliente_id, tipo, descricao, horas_pacote, valor_hora, data_inicio)
                          VALUES (%s, 'banco_horas', 'Zerado', 1, 100, '2026-01-01') RETURNING id""", [self.ids["alfa"]])[0]["id"]
        self.sql("""INSERT INTO apontamentos (usuario_id, contrato_id, data, minutos, descricao, tipo_servico_id)
                    VALUES (%s, %s, current_date, 60, 'gastou', %s)""", [self.ids["tec"], sem, self.tipo("Help desk")])
        r = self.post(tec, f"/chamados/{num}/atualizar", {"contrato_id": str(sem), "responsavel_id": str(self.ids["tec"]),
                                                         "status": "em_atendimento", "prioridade": "media"},
                      follow_redirects=True)
        self.assertIn("sem saldo", r.get_data(as_text=True))
        self.assertIsNone(self.sql("SELECT contrato_id FROM chamados WHERE id = %s", [cid["id"]])[0]["contrato_id"])
        # contrato válido + responsável liberam o atendimento
        bom = self.sql("SELECT id FROM contratos WHERE descricao = 'Segundo Alfa'")[0]["id"]
        self.post(tec, f"/chamados/{num}/atualizar", {"contrato_id": str(bom), "responsavel_id": str(self.ids["tec"]),
                                                     "status": "em_atendimento", "prioridade": "media"})
        self.assertNotIn("Atendimento bloqueado", tec.get(f"/chamados/{num}").get_data(as_text=True))
        self.post(tec, f"/chamados/{num}/responder", {"texto": "agora vai"})
        self.assertTrue(self.sql("SELECT 1 FROM chamado_interacoes WHERE chamado_id = %s", [cid["id"]]))

    def test_cronometros_concorrentes(self):
        ct = self.sql("""INSERT INTO contratos (cliente_id, tipo, descricao, horas_pacote, valor_hora, data_inicio)
                         VALUES (%s, 'banco_horas', 'Pool 2h', 2, 100, '2026-01-01') RETURNING id""",
                      [self.ids["beta"]])[0]["id"]
        nums = []
        for t in ("Pool A", "Pool B"):
            nums.append(self.sql("""INSERT INTO chamados (cliente_id, solicitante_id, contrato_id, responsavel_id, titulo,
                                                          descricao, prioridade, origem)
                                    VALUES (%s, %s, %s, %s, %s, 'x', 'media', 'portal') RETURNING id, numero""",
                                 [self.ids["beta"], self.ids["joao"], ct, self.ids["tec"], t])[0])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, f"/chamados/{nums[1]['numero']}/cronometro/iniciar", {})
        self.sql("UPDATE cronometros SET rodando_desde = now() - interval '30 minutes' WHERE chamado_id = %s",
                 [nums[1]["id"]])
        e = tec.get(f"/chamados/{nums[0]['numero']}/cronometro/estado").get_json()
        self.assertEqual(e["outros_rodando"], 1)
        self.assertTrue(5300 <= e["limite"] <= 5400)  # 2h − 30 min do outro recurso
        # o outro estoura o saldo: este fica sem limite disponível
        self.sql("UPDATE cronometros SET rodando_desde = now() - interval '3 hours' WHERE chamado_id = %s",
                 [nums[1]["id"]])
        self.assertEqual(tec.get(f"/chamados/{nums[0]['numero']}/cronometro/estado").get_json()["limite"], 0)
        self.post(tec, f"/chamados/{nums[0]['numero']}/cronometro/iniciar", {})
        self.assertFalse(self.sql("SELECT 1 FROM cronometros WHERE chamado_id = %s", [nums[0]["id"]]))

    def test_emails_de_aprovacao(self):
        cid = self.sql("SELECT id FROM chamados ORDER BY id LIMIT 1")[0]["id"]
        ids = [self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                           VALUES (%s, %s, current_date, 45, %s, %s) RETURNING id""",
                        [self.ids["tec"], cid, f"trabalho {i}", self.tipo("Help desk")])[0]["id"] for i in range(2)]
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, "/admin/aprovacao", {"ids": [str(i) for i in ids], "acao": "rejeitado", "motivo": "Sem detalhe"})
        emails = self.sql("SELECT para, assunto, corpo FROM emails_saida WHERE evento = 'aprovacao_horas' ORDER BY id")
        para = {e["para"]: e for e in emails}
        self.assertIn("ana@portosolucoes.com.br", para)
        self.assertIn("contato@portosolucoes.com.br", para)
        corpo = para["ana@portosolucoes.com.br"]["corpo"]
        for trecho in ("rejeitados", "trabalho 0", "trabalho 1", "Sem detalhe", "Total: 1h30"):
            self.assertIn(trecho, corpo)

    def test_filtros_e_exportacoes(self):
        cid, num = [self.sql("SELECT id, numero FROM chamados ORDER BY id LIMIT 1")[0][k] for k in ("id", "numero")]
        self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                    VALUES (%s, %s, current_date, 50, 'Troca do switch do andar', %s)""",
                 [self.ids["tec"], cid, self.tipo("Help desk")])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        html = tec.get("/horas?descricao=switch").get_data(as_text=True)
        self.assertIn("Troca do switch", html)
        self.assertNotIn("Troca do switch", tec.get("/horas?descricao=impressora").get_data(as_text=True))
        self.assertIn("Troca do switch", tec.get(f"/horas?chamado={num}").get_data(as_text=True))
        self.assertIn("Troca do switch", tec.get("/horas?duracao=0:50").get_data(as_text=True))
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertIn("Troca do switch", adm.get("/admin/apontamentos?descricao=swit").get_data(as_text=True))
        self.assertIn("Troca do switch", adm.get(f"/admin/aprovacao?chamado={num}").get_data(as_text=True))
        sol = self.sql("SELECT solicitante_id FROM chamados WHERE id = %s", [cid])[0]["solicitante_id"]
        self.assertIn("Troca do switch", adm.get(f"/admin/aprovacao?solicitante={sol}").get_data(as_text=True))
        self.assertNotIn("Troca do switch", adm.get("/admin/aprovacao?status_chamado=fechado").get_data(as_text=True))
        for c, url in [(tec, "/horas"), (adm, "/admin/apontamentos"), (adm, "/admin/aprovacao"), (adm, "/admin/saldos"),
                       (adm, f"/admin/saldos?contrato={self.ids['ct_beta']}")]:
            with self.subTest(url=url):
                sep = "&" if "?" in url else "?"
                self.assertTrue(c.get(url + sep + "formato=pdf").data.startswith(b"%PDF"))
                r = c.get(url + sep + "formato=xlsx")
                self.assertEqual(r.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_tela_do_chamado(self):
        num = self.sql("SELECT numero FROM chamados WHERE contrato_id = %s ORDER BY id LIMIT 1", [self.ids["ct_beta"]])[0]["numero"]
        self.sql("UPDATE chamados SET responsavel_id = %s WHERE numero = %s", [self.ids["tec"], num])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        html = tec.get(f"/chamados/{num}").get_data(as_text=True)
        # cronômetro acima de Detalhes, Atualizar acima de Agenda; detalhes mostram status/prioridade/responsável
        self.assertLess(html.index('id="cronometro"'), html.index("<h3>Detalhes</h3>"))
        self.assertLess(html.index('id="form-atualizar"'), html.index('id="agenda"'))
        for k in ("status", "prioridade", "responsavel_id", "contrato_id"):
            self.assertIn(f'data-espelho="{k}"', html)
        self.assertIn("Ana Técnica", html.split('data-espelho="responsavel_id"')[1][:40])


class TestVersao16(Base):
    """Visibilidade por responsável, avisos de saldo, horas extras, respostas prontas, base de conhecimento e
    fechamento de período."""

    def tipo(self):
        return self.sql("SELECT id FROM tipos_servico WHERE nome = 'Help desk'")[0]["id"]

    def chamado(self, titulo, resp=None, contrato=None):
        return self.sql("""INSERT INTO chamados (cliente_id, solicitante_id, contrato_id, responsavel_id, titulo, descricao,
                                                 prioridade, origem)
                           VALUES (%s, %s, %s, %s, %s, 'x', 'media', 'portal') RETURNING id, numero""",
                        [self.ids["beta"], self.ids["joao"], contrato or self.ids["ct_beta"], resp, titulo])[0]

    def test_versao_no_menu(self):
        from app import VERSAO
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertIn(f"SisPorto v{VERSAO}", adm.get("/painel").get_data(as_text=True))

    def test_visibilidade_por_responsavel(self):
        livre = self.chamado("Visível a todos")
        meu = self.chamado("Da Ana", resp=self.ids["tec"])
        outro = self.chamado("Do admin", resp=self.ids["adm"])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        lista = tec.get("/chamados?situacao=todos").get_data(as_text=True)
        self.assertIn("Visível a todos", lista)
        self.assertIn("Da Ana", lista)
        self.assertNotIn("Do admin", lista)
        self.assertEqual(tec.get(f"/chamados/{livre['numero']}").status_code, 200)
        self.assertEqual(tec.get(f"/chamados/{meu['numero']}").status_code, 200)
        self.assertEqual(tec.get(f"/chamados/{outro['numero']}").status_code, 403)
        self.assertEqual(self.post(tec, f"/chamados/{outro['numero']}/responder", {"texto": "x"}).status_code, 403)
        self.assertEqual(tec.get(f"/chamados/{outro['numero']}/cronometro/estado").status_code, 403)
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertIn("Do admin", adm.get("/chamados?situacao=todos").get_data(as_text=True))
        self.assertIn("Da Ana", adm.get("/chamados?situacao=todos").get_data(as_text=True))

    def test_avisos_de_saldo(self):
        ct = self.sql("""INSERT INTO contratos (cliente_id, tipo, descricao, horas_pacote, valor_hora, data_inicio)
                         VALUES (%s, 'banco_horas', 'Aviso 10h', 10, 100, '2026-01-01') RETURNING id""",
                      [self.ids["beta"]])[0]["id"]
        self.sql("UPDATE usuarios SET recebe_alerta_saldo = true WHERE id = %s", [self.ids["joao"]])
        c = self.chamado("Consumo", resp=self.ids["tec"], contrato=ct)
        tec = self.cliente_http("ana@portosolucoes.com.br")
        hoje = self.sql("SELECT current_date AS d")[0]["d"].isoformat()

        def lancar(d):
            self.post(tec, f"/chamados/{c['numero']}/horas", {"data": hoje, "duracao": d, "descricao": "trabalho",
                                                              "tipo_servico_id": str(self.tipo()), "faturavel": "1"})

        def avisos():
            return self.sql("""SELECT para, assunto FROM emails_saida WHERE evento = 'alerta_saldo'
                                AND assunto LIKE %s ORDER BY id""", ["%Aviso 10h%"])
        lancar("7h")
        self.assertEqual(avisos(), [])
        lancar("1h30")  # 85%
        a = avisos()
        self.assertTrue(any(x["para"] == "joao@clinicabeta.com.br" and "80%" in x["assunto"] for x in a))
        self.assertTrue(any("@portosolucoes" in x["para"] for x in a))
        n = len(a)
        lancar("0:10")  # continua em 80%: não repete
        self.assertEqual(len(avisos()), n)
        self.post(self.cliente_http("contato@portosolucoes.com.br"), f"/chamados/{c['numero']}/horas",
                  {"data": hoje, "duracao": "1:20", "descricao": "fim", "tipo_servico_id": str(self.tipo()), "faturavel": "1"})
        self.assertTrue(any("100%" in x["assunto"] for x in avisos()))

    def test_horas_extras(self):
        ct = self.sql("""INSERT INTO contratos (cliente_id, tipo, descricao, horas_pacote, valor_hora, data_inicio)
                         VALUES (%s, 'banco_horas', 'Extras 2h', 2, 150, '2026-01-01') RETURNING id""",
                      [self.ids["beta"]])[0]["id"]
        c = self.chamado("Precisa de mais", resp=self.ids["tec"], contrato=ct)
        self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                    VALUES (%s, %s, current_date, 120, 'tudo', %s)""", [self.ids["tec"], c["id"], self.tipo()])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertIn("solicitar horas extras", tec.get(f"/chamados/{c['numero']}").get_data(as_text=True).lower())
        r = self.post(tec, "/horas-extras", {"chamado": str(c["numero"]), "origem": "chamado", "horas": "5",
                                             "justificativa": "Migração exige mais 5 horas."})
        self.assertIn(f"/chamados/{c['numero']}", r.headers["Location"])
        e = self.sql("SELECT * FROM horas_extras WHERE contrato_id = %s", [ct])[0]
        self.assertEqual((float(e["horas"]), e["status"], float(e["valor_hora"])), (5, "pendente", 150))
        self.assertTrue(self.sql("SELECT 1 FROM emails_saida WHERE evento = 'horas_extras' AND para = 'joao@clinicabeta.com.br'"))
        # outro cliente não vê; usuário sem marca de aprovador não aprova quando há aprovador marcado
        maria = self.cliente_http("maria@alfalog.com.br")
        self.assertNotIn("Migração", maria.get("/portal/horas-extras").get_data(as_text=True))
        self.assertEqual(self.post(maria, f"/portal/horas-extras/{e['id']}", {"acao": "aprovado"}).status_code, 404)
        joao = self.cliente_http("joao@clinicabeta.com.br")
        self.assertIn("Migração exige", joao.get("/portal/horas-extras").get_data(as_text=True))
        self.post(joao, f"/portal/horas-extras/{e['id']}", {"acao": "recusado"})  # recusa sem motivo: bloqueia
        self.assertEqual(self.sql("SELECT status FROM horas_extras WHERE id = %s", [e["id"]])[0]["status"], "pendente")
        self.post(joao, f"/portal/horas-extras/{e['id']}", {"acao": "aprovado", "resposta": "ok"})
        self.assertEqual(self.sql("SELECT status FROM horas_extras WHERE id = %s", [e["id"]])[0]["status"], "aprovado")
        s = self.sql("SELECT horas_saldo, horas_extras, horas_liberadas FROM vw_saldo_banco_horas WHERE contrato_id = %s",
                     [ct])[0]
        self.assertEqual((float(s["horas_saldo"]), float(s["horas_extras"]), float(s["horas_liberadas"])), (5, 5, 7))
        self.assertTrue(self.sql("SELECT 1 FROM emails_saida WHERE evento = 'horas_extras' AND para = 'ana@portosolucoes.com.br'"))
        self.assertNotIn("Atendimento bloqueado", tec.get(f"/chamados/{c['numero']}").get_data(as_text=True))

    def test_respostas_prontas_e_conhecimento(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.post(adm, "/config/respostas-prontas", {"titulo": "Acesso remoto",
                                                    "texto": "Olá, {solicitante}. No chamado #{numero} precisamos de acesso."})
        c = self.chamado("Impressora não imprime nada", resp=self.ids["tec"])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.post(tec, "/conhecimento/novo", {"titulo": "Impressora não imprime: limpar fila do spooler",
                                              "conteudo": "Abra services.msc e reinicie o Spooler.", "publico": "1",
                                              "categoria": "Impressoras", "palavras": "spooler fila"})
        self.post(tec, "/conhecimento/novo", {"titulo": "Senha do roteador da Porto", "conteudo": "interno", })
        html = tec.get(f"/chamados/{c['numero']}").get_data(as_text=True)
        self.assertIn(f"No chamado #{c['numero']} precisamos de acesso", html)
        self.assertIn("João Lima", html.split("Acesso remoto")[0][-400:] + html)  # {solicitante} preenchido
        self.assertIn("limpar fila do spooler", html)  # artigo relacionado
        joao = self.cliente_http("joao@clinicabeta.com.br")
        sug = joao.get("/conhecimento/buscar.json?q=impressora não imprime").get_json()
        self.assertEqual(len(sug), 1)
        self.assertEqual(joao.get("/conhecimento/buscar.json?q=senha roteador").get_json(), [])
        ajuda = joao.get("/portal/ajuda").get_data(as_text=True)
        self.assertIn("spooler", ajuda)
        self.assertNotIn("roteador", ajuda)
        interno = self.sql("SELECT id FROM artigos WHERE titulo LIKE 'Senha%%'")[0]["id"]
        self.assertEqual(joao.get(f"/conhecimento/{interno}").status_code, 404)
        pub = self.sql("SELECT id FROM artigos WHERE publico")[0]["id"]
        self.assertEqual(joao.get(f"/conhecimento/{pub}").status_code, 200)
        self.post(joao, f"/conhecimento/{pub}/util", {"util": "1"})
        self.assertEqual(self.sql("SELECT util_sim, visualizacoes FROM artigos WHERE id = %s", [pub])[0]["util_sim"], 1)
        self.assertEqual(self.cliente_http("joao@clinicabeta.com.br").get("/conhecimento").status_code, 403)

    def test_fechamento_de_periodo(self):
        cid = self.sql("SELECT id FROM chamados ORDER BY id LIMIT 1")[0]["id"]
        ap = self.sql("""INSERT INTO apontamentos (usuario_id, chamado_id, data, minutos, descricao, tipo_servico_id)
                         VALUES (%s, %s, '2026-03-10', 60, 'março', %s) RETURNING id""",
                      [self.ids["tec"], cid, self.tipo()])[0]["id"]
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = self.post(adm, "/admin/fechamento", {"acao": "fechar", "de": "2026-03-01", "ate": "2026-03-31"},
                      follow_redirects=True)
        self.assertIn("pendente", r.get_data(as_text=True))
        self.assertFalse(self.sql("SELECT 1 FROM periodos_fechados"))
        self.post(adm, "/admin/aprovacao", {"ids": str(ap), "acao": "aprovado"})
        self.post(adm, "/admin/fechamento", {"acao": "fechar", "de": "2026-03-01", "ate": "2026-03-31"})
        self.assertTrue(self.sql("SELECT 1 FROM periodos_fechados"))
        # banco recusa qualquer alteração
        with self.assertRaises(Exception):
            self.sql("UPDATE apontamentos SET minutos = 30 WHERE id = %s", [ap])
        # telas recusam com aviso
        tec = self.cliente_http("ana@portosolucoes.com.br")
        r = self.post(tec, "/horas/lancar", {"contrato_id": str(self.ids["ct_beta"]), "data": "2026-03-15",
                                             "duracao": "1h", "descricao": "atrasado", "tipo_servico_id": str(self.tipo())},
                      follow_redirects=True)
        self.assertIn("está fechado", r.get_data(as_text=True))
        r = self.post(adm, f"/admin/apontamentos/{ap}", {"usuario_id": str(self.ids["tec"]), "chamado": "",
                                                        "contrato_id": str(self.ids["ct_beta"]), "status": "aprovado",
                                                        "data": "2026-03-10", "duracao": "2h", "descricao": "x",
                                                        "tipo_servico_id": str(self.tipo())}, follow_redirects=True)
        self.assertIn("período fechado", r.get_data(as_text=True))
        self.assertEqual(self.sql("SELECT minutos FROM apontamentos WHERE id = %s", [ap])[0]["minutos"], 60)
        # reabrir libera
        pid = self.sql("SELECT id FROM periodos_fechados")[0]["id"]
        self.post(adm, "/admin/fechamento", {"acao": "reabrir", "id": str(pid)})
        self.sql("UPDATE apontamentos SET minutos = 30 WHERE id = %s", [ap])
        self.assertEqual(self.sql("SELECT minutos FROM apontamentos WHERE id = %s", [ap])[0]["minutos"], 30)
        self.assertEqual(self.cliente_http("ana@portosolucoes.com.br").get("/admin/fechamento").status_code, 403)


class TestPacote(unittest.TestCase):
    def test_rotas_e_telas(self):
        from app.verificar import verificar
        self.assertEqual(verificar(), [])

    def test_detecta_rota_faltando(self):
        """Simula o problema da 1.6.0: tela nova com módulo .py antigo (rota inexistente)."""
        import pathlib, tempfile
        from app import create_app
        from app.verificar import verificar
        with tempfile.TemporaryDirectory() as d:
            alvo = pathlib.Path(__file__).resolve().parent.parent / "app" / "templates" / "_teste_rota.html"
            alvo.write_text("{{ url_for('kb.nao_existe') }}", encoding="utf-8")
            try:
                erros = verificar()
            finally:
                alvo.unlink()
        self.assertTrue(any("kb.nao_existe" in e for e in erros))


class TestVersao17(Base):
    """Ajustes 02 + limpeza automática + limite de tentativas."""

    def test_cliente_dados_no_topo_e_valores_ocultos(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        html = adm.get(f"/clientes/{self.ids['alfa']}").get_data(as_text=True)
        self.assertLess(html.index("Dados do cliente"), html.index("<h2>Contratos</h2>"))
        self.assertIn("R$", html.split("<h2>Contratos</h2>")[1].split("Usuários do cliente")[0])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        html = tec.get(f"/clientes/{self.ids['alfa']}").get_data(as_text=True)
        self.assertLess(html.index("Dados do cliente"), html.index("<h2>Contratos</h2>"))
        self.assertNotIn("R$", html.split("<h2>Contratos</h2>")[1].split("Usuários do cliente")[0])

    def test_tecnico_altera_status_prioridade_responsavel(self):
        num = self.sql("""INSERT INTO chamados (cliente_id, solicitante_id, contrato_id, titulo, descricao, prioridade, origem)
                          VALUES (%s, %s, %s, 'Status pelo técnico', 'x', 'media', 'portal') RETURNING numero""",
                       [self.ids["beta"], self.ids["joao"], self.ids["ct_beta"]])[0]["numero"]
        self.sql("INSERT INTO permissoes (usuario_id, rotina, permitido) VALUES (%s, 'chamados_gerir', false)", [self.ids["tec"]])
        tec = self.cliente_http("ana@portosolucoes.com.br")
        self.assertIn('id="form-atualizar"', tec.get(f"/chamados/{num}").get_data(as_text=True))
        self.post(tec, f"/chamados/{num}/atualizar", {"status": "em_atendimento", "prioridade": "alta",
                                                     "responsavel_id": str(self.ids["tec"])})
        c = self.sql("SELECT status::text, prioridade::text, responsavel_id, contrato_id FROM chamados WHERE numero = %s", [num])[0]
        self.assertEqual((c["status"], c["prioridade"], c["responsavel_id"], c["contrato_id"]),
                         ("em_atendimento", "alta", self.ids["tec"], self.ids["ct_beta"]))
        self.sql("DELETE FROM permissoes WHERE usuario_id = %s", [self.ids["tec"]])

    def test_limite_de_tentativas_no_login(self):
        c = self.app.test_client()
        c.get("/entrar")
        for _ in range(8):
            self.post(c, "/entrar", {"email": "ana@portosolucoes.com.br", "senha": "errada"})
        r = self.post(c, "/entrar", {"email": "ana@portosolucoes.com.br", "senha": SENHA_DEMO})
        self.assertEqual(r.status_code, 429)
        self.sql("DELETE FROM tentativas")

    def test_limpeza_automatica(self):
        cid = self.sql("SELECT id FROM chamados ORDER BY id LIMIT 1")[0]["id"]
        aid = self.sql("""INSERT INTO chamado_anexos (chamado_id, nome_arquivo, tamanho_bytes, tipo_mime, arquivo, drive_arquivo_id,
                                                      drive_url, criado_em)
                          VALUES (%s, 'velho.zip', 3, 'application/zip', '\\x504b03', 'x', 'https://drive.google.com/x',
                                  now() - interval '120 days') RETURNING id""", [cid])[0]["id"]
        self.sql("""INSERT INTO emails_saida (para, assunto, corpo, evento, enviado_em)
                    VALUES ('a@b.com', 'velho', 'x', 'teste', now() - interval '200 days')""")
        from app.limpeza import limpar
        with self.app.app_context():
            from app.db import db
            r = limpar(forcar=True)
            db().commit()
        self.assertGreaterEqual(r["anexos_arquivados"], 1)
        self.assertIsNone(self.sql("SELECT arquivo FROM chamado_anexos WHERE id = %s", [aid])[0]["arquivo"])
        self.assertFalse(self.sql("SELECT 1 FROM emails_saida WHERE assunto = 'velho'"))
        adm = self.cliente_http("contato@portosolucoes.com.br")
        r = adm.get(f"/anexos/{aid}")
        self.assertEqual(r.headers.get("Location"), "https://drive.google.com/x")


class TestDominio(Base):
    """1.7.1: sistema atrás do Firebase Hosting (sistema.portosolucoes.com.br)."""

    def test_cookie_de_sessao_compativel_com_firebase(self):
        c = self.app.test_client()
        r = c.get("/entrar")
        self.assertIn("__session=", r.headers.get("Set-Cookie", ""))
        self.cliente_http("contato@portosolucoes.com.br")  # login continua funcionando

    def test_paginas_nunca_ficam_no_cache_do_cdn(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        self.assertEqual(adm.get("/painel").headers["Cache-Control"], "private, no-store")
        self.assertEqual(adm.get("/entrar").headers["Cache-Control"], "private, no-store")
        self.assertIn("public", self.app.test_client().get("/marca/logo").headers.get("Cache-Control", ""))

    def test_links_usam_o_dominio_oficial(self):
        from flask import url_for
        app = create_app({"DATABASE_URL": TEST_DSN, "TESTING": True, "SECRET_KEY": "teste",
                          "URL_PUBLICA": "https://sistema.portosolucoes.com.br"})
        c = app.test_client()
        r = c.get("/", headers={"Host": "sisporto-xyz-rj.a.run.app"})
        self.assertTrue(r.headers["Location"].startswith(("https://sistema.portosolucoes.com.br/", "/entrar")))
        with app.test_request_context("/", base_url="https://sistema.portosolucoes.com.br"):
            self.assertEqual(url_for("auth.google_retorno", _external=True),
                             "https://sistema.portosolucoes.com.br/entrar/google/retorno")

    def test_ip_real_atras_do_firebase(self):
        from app.auth import ip_cliente
        app = create_app({"DATABASE_URL": TEST_DSN, "TESTING": True, "SECRET_KEY": "teste",
                          "URL_PUBLICA": "https://sistema.portosolucoes.com.br"})
        with app.test_request_context("/", headers={"X-Forwarded-For": "200.1.2.3, 35.4.5.6",
                                                    "X-Forwarded-Host": "sistema.portosolucoes.com.br"},
                                      environ_base={"REMOTE_ADDR": "35.4.5.6"}):
            self.assertEqual(ip_cliente(), "200.1.2.3")
        with app.test_request_context("/", headers={"X-Forwarded-For": "200.1.2.3, 35.4.5.6"},
                                      environ_base={"REMOTE_ADDR": "35.4.5.6"}):
            self.assertEqual(ip_cliente(), "35.4.5.6")

    def test_orcamento_aprovado(self):
        conf = open(os.path.join(os.path.dirname(__file__), "..", "infra", "gcp", "orcamento.conf")).read()
        self.assertIn("TETO_BRL=400", conf)
        self.assertIn("CORTE_BRL=500", conf)
        self.assertIn('EMAIL_NOTIFICACOES="contato@portosolucoes.com.br"', conf)


class TestIdentidade18(Base):
    """1.8.0: nova identidade visual do SisPorto."""

    def test_marca_e_fontes_servidas(self):
        c = self.app.test_client()
        html = c.get("/entrar").get_data(as_text=True)
        self.assertIn("marca/simbolo-sisdark.svg", html)
        self.assertIn("marca/favicon.svg", html)
        self.assertIn("BY PORTO SOLUÇÕES", html)
        for arq in ("marca/simbolo-sis.svg", "marca/simbolo-sisdark.svg", "marca/favicon.svg", "marca/icone-180.png",
                    "fontes/inter-regular.woff", "fontes/inter-bold.woff"):
            r = c.get(f"/static/{arq}")
            self.assertEqual(r.status_code, 200, arq)
            r.close()
        css = c.get("/static/app.css").get_data(as_text=True)
        self.assertIn("--turquesa:#22C0B2", css)

    def test_menu_sem_emoji(self):
        adm = self.cliente_http("contato@portosolucoes.com.br")
        html = adm.get("/painel").get_data(as_text=True)
        self.assertIn('class="marca-logo"', html)
        for e in ("📊", "🎧", "⏱", "📄", "🗂", "⚙", "🛡"):
            self.assertNotIn(e, html)
