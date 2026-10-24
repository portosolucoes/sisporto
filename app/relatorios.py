"""Relatório de apontamentos de horas (por cliente, período, recurso e tipo de serviço)."""
from __future__ import annotations

import datetime as dt
import io

from flask import Blueprint, Response, render_template, request

from .acessos import requer
from .db import db
from .util import TZ

bp = Blueprint("relatorios", __name__)


def _int(v):
    try:
        return int(v) if v else None
    except ValueError:
        return None


def _filtros():
    hoje = dt.datetime.now(TZ).date()
    f = {k: request.args.get(k, "") for k in ("de", "ate", "cliente", "recurso", "tipo", "status", "faturavel")}
    try:
        de = dt.date.fromisoformat(f["de"]) if f["de"] else hoje.replace(day=1)
        ate = dt.date.fromisoformat(f["ate"]) if f["ate"] else hoje
    except ValueError:
        de, ate = hoje.replace(day=1), hoje
    f["de"], f["ate"] = de.isoformat(), ate.isoformat()
    where = ["a.data BETWEEN %s AND %s"]
    params = [de, ate]
    if _int(f["cliente"]):
        where.append("cl.id = %s"); params.append(int(f["cliente"]))
    if _int(f["recurso"]):
        where.append("a.usuario_id = %s"); params.append(int(f["recurso"]))
    if _int(f["tipo"]):
        where.append("a.tipo_servico_id = %s"); params.append(int(f["tipo"]))
    if f["status"] in ("pendente", "aprovado", "rejeitado"):
        where.append("a.status = %s"); params.append(f["status"])
    elif f["status"] != "todos":
        where.append("a.status <> 'rejeitado'")
    if f["faturavel"] in ("sim", "nao"):
        where.append("a.faturavel = %s"); params.append(f["faturavel"] == "sim")
    return f, " AND ".join(where), params


BASE = """
  FROM apontamentos a
  JOIN usuarios u ON u.id = a.usuario_id
  LEFT JOIN chamados c ON c.id = a.chamado_id
  LEFT JOIN contratos ct ON ct.id = coalesce(a.contrato_id, c.contrato_id)
  LEFT JOIN clientes cl ON cl.id = coalesce(c.cliente_id, ct.cliente_id)
  LEFT JOIN tipos_servico ts ON ts.id = a.tipo_servico_id
"""


def _dados():
    f, where, params = _filtros()
    linhas = db().all(f"""SELECT a.data, a.minutos, a.faturavel, a.status, a.descricao, u.nome AS recurso,
                                 u.codigo AS recurso_codigo, coalesce(cl.nome_fantasia, cl.razao_social) AS cliente,
                                 ct.descricao AS contrato, c.numero, c.titulo, coalesce(ts.nome, 'Sem tipo') AS tipo
                          {BASE} WHERE {where} ORDER BY a.data, cliente, a.id LIMIT 5000""", params)
    tot = db().one(f"""SELECT coalesce(sum(a.minutos),0) AS total,
                              coalesce(sum(a.minutos) FILTER (WHERE a.faturavel),0) AS faturavel,
                              coalesce(sum(a.minutos) FILTER (WHERE a.status = 'aprovado'),0) AS aprovado,
                              coalesce(sum(a.minutos) FILTER (WHERE a.status = 'pendente'),0) AS pendente
                         {BASE} WHERE {where}""", params)

    def agrupar(expr):
        return db().all(f"""SELECT {expr} AS nome, sum(a.minutos) AS minutos,
                                   sum(a.minutos) FILTER (WHERE a.faturavel) AS faturavel
                              {BASE} WHERE {where} GROUP BY 1 ORDER BY 2 DESC""", params)

    grupos = {
        "Por cliente": agrupar("coalesce(cl.nome_fantasia, cl.razao_social, '—')"),
        "Por tipo de serviço": agrupar("coalesce(ts.nome, 'Sem tipo')"),
        "Por recurso": agrupar("u.nome"),
    }
    return f, linhas, tot, grupos


@bp.get("/apontamentos")
@requer("relatorios")
def apontamentos():
    f, linhas, tot, grupos = _dados()
    if request.args.get("formato") == "xlsx":
        return _excel(f, linhas, grupos)
    clientes = db().all("SELECT id, coalesce(nome_fantasia, razao_social) AS nome FROM clientes ORDER BY 2")
    recursos = db().all("SELECT id, nome FROM usuarios WHERE perfil <> 'cliente' ORDER BY nome")
    tipos = db().all("SELECT id, nome FROM tipos_servico ORDER BY ordem, nome")
    return render_template("equipe/relatorio_apontamentos.html", f=f, linhas=linhas, tot=tot, grupos=grupos,
                           clientes=clientes, recursos=recursos, tipos=tipos)


def _excel(f, linhas, grupos):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Apontamentos"
    cab = ["Data", "Cliente", "Contrato", "Chamado", "Título", "Tipo de serviço", "Recurso", "Código",
           "Descrição", "Horas", "Faturável", "Status"]
    ws.append(cab)
    for l in linhas:
        ws.append([l["data"], l["cliente"], l["contrato"], l["numero"], l["titulo"], l["tipo"], l["recurso"],
                   l["recurso_codigo"], l["descricao"], round(l["minutos"] / 60, 2),
                   "Sim" if l["faturavel"] else "Não", l["status"]])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for col, w in zip("ABCDEFGHIJKL", (11, 26, 24, 9, 30, 18, 20, 9, 45, 8, 10, 11)):
        ws.column_dimensions[col].width = w
    for r in range(2, ws.max_row + 1):
        ws[f"A{r}"].number_format = "DD/MM/YYYY"
        ws[f"J{r}"].number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    res = wb.create_sheet("Resumo")
    res.append([f"Período: {f['de']} a {f['ate']}"])
    for titulo, itens in grupos.items():
        res.append([])
        res.append([titulo, "Horas", "Horas faturáveis"])
        res.cell(row=res.max_row, column=1).font = Font(bold=True)
        for i in itens:
            res.append([i["nome"], round(i["minutos"] / 60, 2), round((i["faturavel"] or 0) / 60, 2)])
    res.column_dimensions["A"].width = 34
    buf = io.BytesIO()
    wb.save(buf)
    nome = f"apontamentos_{f['de']}_{f['ate']}.xlsx"
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={nome}"})


# =============================================================================
# Central de relatórios (modelos de mercado para consultorias de TI)
# =============================================================================
from flask import abort, url_for  # noqa: E402

from .acessos import pode  # noqa: E402

CATALOGO = [
    # chave, título, descrição, rotina exigida
    ("apontamentos", "Apontamentos de horas", "Horas por cliente, recurso e tipo de serviço, com detalhe dos lançamentos.", "relatorios"),
    ("sla", "SLA e desempenho do atendimento", "% de chamados no prazo, tempo médio de 1ª resposta e de solução, por prioridade e cliente.", "relatorios"),
    ("chamados", "Chamados por período", "Volume aberto e encerrado, por status, categoria, origem e cliente; envelhecimento do backlog.", "relatorios"),
    ("produtividade", "Produtividade por recurso", "Horas, % faturável, chamados resolvidos e avaliação média por recurso.", "relatorios"),
    ("contratos", "Consumo de contratos", "Horas liberadas x consumidas, saldo, franquia mensal e excedentes.", "relatorios"),
    ("satisfacao", "Satisfação (CSAT e NPS)", "Avaliação dos chamados (1 a 5) e Net Promoter Score mensal dos clientes.", "relatorios"),
    ("prefaturamento", "Pré-faturamento do mês", "Quanto faturar de cada cliente: mensalidades, excedentes e horas avulsas aprovadas.", "relatorios_financeiros"),
    ("rentabilidade", "Rentabilidade por cliente e contrato", "Receita estimada x custo dos recursos (custo/hora) e margem.", "relatorios_financeiros"),
]


def _periodo():
    hoje = dt.datetime.now(TZ).date()
    try:
        de = dt.date.fromisoformat(request.args["de"]) if request.args.get("de") else hoje.replace(day=1)
        ate = dt.date.fromisoformat(request.args["ate"]) if request.args.get("ate") else hoje
    except ValueError:
        de, ate = hoje.replace(day=1), hoje
    cliente = _int(request.args.get("cliente"))
    return de, ate, cliente


def _fcli(alias="cl", cliente=None):
    return (f" AND {alias}.id = %s", [cliente]) if cliente else ("", [])


def r_sla(de, ate, cliente):
    fc, pc = _fcli("cl", cliente)
    base = f"""FROM chamados c JOIN clientes cl ON cl.id = c.cliente_id
               WHERE c.criado_em::date BETWEEN %s AND %s {fc}"""
    p = [de, ate] + pc
    met = """count(*) AS total,
             count(*) FILTER (WHERE c.resolvido_em IS NOT NULL) AS resolvidos,
             count(*) FILTER (WHERE c.resolvido_em IS NOT NULL AND (c.prazo_solucao IS NULL OR c.resolvido_em <= c.prazo_solucao)) AS no_prazo,
             round(avg(fn_minutos_uteis_entre(c.criado_em, c.primeira_resposta_em)) FILTER (WHERE c.primeira_resposta_em IS NOT NULL)) AS media_resposta,
             round(avg(greatest(fn_minutos_uteis_entre(c.criado_em, c.resolvido_em) - c.minutos_pausados, 0)) FILTER (WHERE c.resolvido_em IS NOT NULL)) AS media_solucao"""
    tot = db().one(f"SELECT {met} {base}", p)
    def pct(r):
        return round(100 * r["no_prazo"] / r["resolvidos"], 1) if r["resolvidos"] else None
    por_prio = db().all(f"""SELECT c.prioridade::text AS grupo, {met} {base} GROUP BY 1
                            ORDER BY array_position(ARRAY['critica','alta','media','baixa'], c.prioridade::text)""", p)
    por_cli = db().all(f"""SELECT coalesce(cl.nome_fantasia, cl.razao_social) AS grupo, {met} {base}
                           GROUP BY 1 ORDER BY 2 DESC""", p)
    for r in por_prio + por_cli:
        r["pct"] = pct(r)
    cols = [("grupo", "Grupo", "texto"), ("total", "Abertos", "num"), ("resolvidos", "Resolvidos", "num"),
            ("no_prazo", "No prazo", "num"), ("pct", "% no prazo", "pct"),
            ("media_resposta", "Média 1ª resposta (útil)", "horas"), ("media_solucao", "Média solução (útil)", "horas")]
    return {"cards": [("Chamados abertos", tot["total"], "num"), ("% no prazo (resolvidos)", pct(tot), "pct"),
                      ("Média 1ª resposta", tot["media_resposta"], "horas"), ("Média de solução", tot["media_solucao"], "horas")],
            "tabelas": [{"titulo": "Por prioridade", "colunas": [("grupo", "Prioridade", "prioridade")] + cols[1:],
                         "linhas": por_prio},
                        {"titulo": "Por cliente", "colunas": [("grupo", "Cliente", "texto")] + cols[1:], "linhas": por_cli}],
            "nota": "Tempos em horário útil, descontando o período em 'aguardando cliente'."}


def r_chamados(de, ate, cliente):
    fc, pc = _fcli("cl", cliente)
    base = f"FROM chamados c JOIN clientes cl ON cl.id = c.cliente_id WHERE c.criado_em::date BETWEEN %s AND %s {fc}"
    p = [de, ate] + pc
    def grp(expr, nome):
        return {"titulo": f"Por {nome}", "colunas": [("grupo", nome.capitalize(), "texto"), ("n", "Chamados", "num")],
                "linhas": db().all(f"SELECT {expr} AS grupo, count(*) AS n {base} GROUP BY 1 ORDER BY 2 DESC", p)}
    tot = db().one(f"""SELECT count(*) AS abertos, count(*) FILTER (WHERE c.status IN ('resolvido','fechado')) AS encerrados
                       {base}""", p)
    encerrados_periodo = db().val(f"""SELECT count(*) FROM chamados c JOIN clientes cl ON cl.id = c.cliente_id
                                      WHERE c.resolvido_em::date BETWEEN %s AND %s {fc}""", p)
    aging = db().all(f"""SELECT CASE WHEN idade <= 2 THEN '0 a 2 dias' WHEN idade <= 7 THEN '3 a 7 dias'
                                     WHEN idade <= 15 THEN '8 a 15 dias' ELSE 'Mais de 15 dias' END AS grupo,
                                count(*) AS n, min(idade) AS ordem
                           FROM (SELECT current_date - c.criado_em::date AS idade FROM chamados c
                                   JOIN clientes cl ON cl.id = c.cliente_id
                                  WHERE c.status IN ('novo','em_atendimento','aguardando_cliente') {fc}) x
                          GROUP BY 1 ORDER BY 3""", pc)
    st = grp("c.status::text", "status")
    for r in st["linhas"]:
        r["grupo"] = {"novo": "Novo", "em_atendimento": "Em atendimento", "aguardando_cliente": "Aguardando cliente",
                      "resolvido": "Resolvido", "fechado": "Fechado"}.get(r["grupo"], r["grupo"])
    org = grp("c.origem::text", "origem")
    for r in org["linhas"]:
        r["grupo"] = {"portal": "Portal", "email": "E-mail", "interno": "Interno"}.get(r["grupo"], r["grupo"])
    backlog = db().val(f"""SELECT count(*) FROM chamados c JOIN clientes cl ON cl.id = c.cliente_id
                           WHERE c.status IN ('novo','em_atendimento','aguardando_cliente') {fc}""", pc)
    return {"cards": [("Abertos no período", tot["abertos"], "num"), ("Resolvidos no período", encerrados_periodo, "num"),
                      ("Backlog atual (em aberto)", backlog, "num"),
                      ("Saldo do período", tot["abertos"] - encerrados_periodo, "num")],
            "tabelas": [st, grp("coalesce(c.categoria, 'Sem categoria')", "categoria"), org,
                        grp("coalesce(cl.nome_fantasia, cl.razao_social)", "cliente"),
                        {"titulo": "Envelhecimento do backlog (hoje)",
                         "colunas": [("grupo", "Idade", "texto"), ("n", "Chamados", "num")], "linhas": aging}]}


def r_produtividade(de, ate, cliente):
    fc, pc = _fcli("cl", cliente)
    linhas = db().all(f"""
        WITH h AS (SELECT a.usuario_id, sum(a.minutos) AS minutos,
                          sum(a.minutos) FILTER (WHERE a.faturavel) AS fat
                     FROM apontamentos a LEFT JOIN contratos ct ON ct.id = a.contrato_id
                     LEFT JOIN clientes cl ON cl.id = ct.cliente_id
                    WHERE a.status <> 'rejeitado' AND a.data BETWEEN %s AND %s {fc} GROUP BY 1),
             r AS (SELECT c.responsavel_id AS usuario_id, count(*) AS resolvidos, round(avg(c.avaliacao), 2) AS csat
                     FROM chamados c JOIN clientes cl ON cl.id = c.cliente_id
                    WHERE c.resolvido_em::date BETWEEN %s AND %s {fc} GROUP BY 1)
        SELECT u.nome AS grupo, u.codigo, coalesce(h.minutos, 0) AS minutos, coalesce(h.fat, 0) AS fat,
               CASE WHEN h.minutos > 0 THEN round(100.0 * h.fat / h.minutos, 1) END AS pct_fat,
               coalesce(r.resolvidos, 0) AS resolvidos, r.csat,
               round(coalesce(h.minutos, 0) / 60.0 * coalesce(u.custo_hora, 0), 2) AS custo
          FROM usuarios u LEFT JOIN h ON h.usuario_id = u.id LEFT JOIN r ON r.usuario_id = u.id
         WHERE u.perfil <> 'cliente' AND (h.minutos IS NOT NULL OR r.resolvidos IS NOT NULL)
         ORDER BY minutos DESC""", [de, ate] + pc + [de, ate] + pc)
    tot_min = sum(l["minutos"] for l in linhas)
    tot_fat = sum(l["fat"] for l in linhas)
    return {"cards": [("Horas lançadas", tot_min, "horas"), ("Horas faturáveis", tot_fat, "horas"),
                      ("% faturável", round(100 * tot_fat / tot_min, 1) if tot_min else None, "pct"),
                      ("Chamados resolvidos", sum(l["resolvidos"] for l in linhas), "num")],
            "tabelas": [{"titulo": "Por recurso", "colunas": [
                ("codigo", "Código", "texto"), ("grupo", "Recurso", "texto"), ("minutos", "Horas", "horas"),
                ("fat", "Faturáveis", "horas"), ("pct_fat", "% faturável", "pct"), ("resolvidos", "Resolvidos", "num"),
                ("csat", "Avaliação média", "num"), ("custo", "Custo (R$)", "moeda")], "linhas": linhas}]}


def r_contratos(de, ate, cliente):
    fc, pc = _fcli("cl", cliente)
    linhas = db().all(f"""
        SELECT coalesce(cl.nome_fantasia, cl.razao_social) AS cliente, ct.descricao, ct.tipo::text AS tipo,
               ct.horas_pacote, ct.franquia_horas, ct.valor_hora,
               coalesce(sum(a.minutos) FILTER (WHERE a.data BETWEEN %s AND %s), 0) AS min_periodo,
               coalesce(sum(a.minutos) FILTER (WHERE a.faturavel AND a.data BETWEEN %s AND %s), 0) AS fat_periodo,
               coalesce(sum(a.minutos) FILTER (WHERE a.faturavel), 0) AS fat_total
          FROM contratos ct JOIN clientes cl ON cl.id = ct.cliente_id
          LEFT JOIN apontamentos a ON a.contrato_id = ct.id AND a.status <> 'rejeitado'
         WHERE ct.ativo {fc}
         GROUP BY ct.id, cl.id ORDER BY cliente, ct.descricao""", [de, ate, de, ate] + pc)
    rot = {"mensal_fixo": "Mensal fixo", "banco_horas": "Banco de horas", "hora_avulsa": "Hora avulsa",
           "projeto": "Projeto", "proposta_avulsa": "Proposta avulsa"}
    for l in linhas:
        l["tipo"] = rot.get(l["tipo"], l["tipo"])
        l["liberadas"] = int(l["horas_pacote"] * 60) if l["horas_pacote"] is not None else None
        l["saldo"] = l["liberadas"] - l["fat_total"] if l["liberadas"] is not None else None
        l["pct"] = round(100 * l["fat_total"] / l["liberadas"], 1) if l["liberadas"] else None
        franq = int(l["franquia_horas"] * 60) if l["franquia_horas"] is not None else None
        l["franquia"] = franq
        l["excedente"] = max(l["fat_periodo"] - franq, 0) if franq is not None else None
        l["valor_exc"] = round(l["excedente"] / 60 * float(l["valor_hora"] or 0), 2) if l["excedente"] else None
    alerta = [l for l in linhas if l["pct"] is not None and l["pct"] >= 80]
    return {"cards": [("Contratos ativos", len(linhas), "num"), ("Com 80%+ das horas usadas", len(alerta), "num"),
                      ("Horas no período", sum(l["min_periodo"] for l in linhas), "horas"),
                      ("Excedente estimado (R$)", sum(l["valor_exc"] or 0 for l in linhas), "moeda")],
            "tabelas": [{"titulo": "Contratos", "colunas": [
                ("cliente", "Cliente", "texto"), ("descricao", "Contrato", "texto"), ("tipo", "Tipo", "texto"),
                ("liberadas", "Horas liberadas", "horas"), ("fat_total", "Consumidas (total)", "horas"),
                ("saldo", "Saldo", "horas"), ("pct", "% usado", "pct"), ("min_periodo", "Horas no período", "horas"),
                ("franquia", "Franquia/mês", "horas"), ("excedente", "Excedente no período", "horas"),
                ("valor_exc", "Excedente (R$)", "moeda")], "linhas": linhas}],
            "nota": "Excedente considera as horas faturáveis do período acima da franquia mensal (ideal: período = 1 mês)."}


def r_satisfacao(de, ate, cliente):
    fc, pc = _fcli("cl", cliente)
    csat = db().one(f"""SELECT count(avaliacao) AS n, round(avg(avaliacao), 2) AS media
                          FROM chamados c JOIN clientes cl ON cl.id = c.cliente_id
                         WHERE c.resolvido_em::date BETWEEN %s AND %s {fc}""", [de, ate] + pc)
    dist = db().all(f"""SELECT c.avaliacao::text || ' estrela(s)' AS grupo, count(*) AS n
                          FROM chamados c JOIN clientes cl ON cl.id = c.cliente_id
                         WHERE c.avaliacao IS NOT NULL AND c.resolvido_em::date BETWEEN %s AND %s {fc}
                         GROUP BY c.avaliacao ORDER BY c.avaliacao DESC""", [de, ate] + pc)
    nps_base = f"""FROM nps_respostas n JOIN clientes cl ON cl.id = n.cliente_id
                  WHERE n.nota IS NOT NULL AND n.respondido_em::date BETWEEN %s AND %s {fc}"""
    met = """count(*) AS respostas,
             count(*) FILTER (WHERE n.nota >= 9) AS promotores,
             count(*) FILTER (WHERE n.nota BETWEEN 7 AND 8) AS neutros,
             count(*) FILTER (WHERE n.nota <= 6) AS detratores"""
    def score(r):
        return round(100 * (r["promotores"] - r["detratores"]) / r["respostas"]) if r["respostas"] else None
    tot = db().one(f"SELECT {met} {nps_base}", [de, ate] + pc)
    por_cli = db().all(f"SELECT coalesce(cl.nome_fantasia, cl.razao_social) AS grupo, {met} {nps_base} GROUP BY 1 ORDER BY 2 DESC",
                       [de, ate] + pc)
    for r in por_cli:
        r["nps"] = score(r)
    enviados = db().val(f"""SELECT count(*) FROM nps_respostas n JOIN clientes cl ON cl.id = n.cliente_id
                            WHERE n.enviado_em::date BETWEEN %s AND %s {fc}""", [de, ate] + pc)
    comentarios = db().all(f"""SELECT n.respondido_em::date AS data, coalesce(cl.nome_fantasia, cl.razao_social) AS cliente,
                                      u.nome AS usuario, n.nota, n.comentario {nps_base.replace('FROM nps_respostas n', 'FROM nps_respostas n JOIN usuarios u ON u.id = n.usuario_id')}
                                  AND n.comentario IS NOT NULL ORDER BY n.respondido_em DESC LIMIT 100""", [de, ate] + pc)
    return {"cards": [("Avaliação média dos chamados (1–5)", csat["media"], "num"),
                      ("Chamados avaliados", csat["n"], "num"), ("NPS", score(tot), "num"),
                      ("Respostas NPS / enviadas", f"{tot['respostas']} / {enviados}", "texto")],
            "tabelas": [{"titulo": "Distribuição das avaliações (CSAT)",
                         "colunas": [("grupo", "Avaliação", "texto"), ("n", "Chamados", "num")], "linhas": dist},
                        {"titulo": "NPS por cliente", "colunas": [
                            ("grupo", "Cliente", "texto"), ("respostas", "Respostas", "num"),
                            ("promotores", "Promotores (9–10)", "num"), ("neutros", "Neutros (7–8)", "num"),
                            ("detratores", "Detratores (0–6)", "num"), ("nps", "NPS", "num")], "linhas": por_cli},
                        {"titulo": "Comentários do NPS", "colunas": [
                            ("data", "Data", "data"), ("cliente", "Cliente", "texto"), ("usuario", "Usuário", "texto"),
                            ("nota", "Nota", "num"), ("comentario", "Comentário", "texto")], "linhas": comentarios}],
            "nota": "NPS = % promotores − % detratores (de −100 a +100)."}


SQL_COBRANCA = """
WITH meses AS (SELECT generate_series(date_trunc('month', %s::date), date_trunc('month', %s::date), interval '1 month')::date AS mes),
h AS (SELECT a.contrato_id, date_trunc('month', a.data)::date AS mes,
             sum(a.minutos) FILTER (WHERE a.faturavel AND a.status = 'aprovado') AS fat_aprov,
             sum(a.minutos) FILTER (WHERE a.faturavel) AS fat,
             sum(a.minutos) FILTER (WHERE a.status = 'pendente') AS pendente,
             sum(a.minutos) AS total,
             sum(a.minutos / 60.0 * coalesce(u.custo_hora, 0)) AS custo
        FROM apontamentos a JOIN usuarios u ON u.id = a.usuario_id
       WHERE a.status <> 'rejeitado' AND a.contrato_id IS NOT NULL GROUP BY 1, 2),
base AS (
  SELECT m.mes, ct.id AS contrato_id, ct.descricao, ct.tipo::text AS tipo, cl.id AS cliente_id,
         coalesce(cl.nome_fantasia, cl.razao_social) AS cliente, ct.valor_mensal, ct.franquia_horas, ct.valor_hora,
         coalesce(h.fat_aprov, 0) AS fat_aprov, coalesce(h.fat, 0) AS fat, coalesce(h.pendente, 0) AS pendente,
         coalesce(h.total, 0) AS total, coalesce(h.custo, 0) AS custo,
         (ct.data_inicio <= (m.mes + interval '1 month - 1 day')::date AND (ct.data_fim IS NULL OR ct.data_fim >= m.mes)) AS vigente
    FROM meses m CROSS JOIN contratos ct JOIN clientes cl ON cl.id = ct.cliente_id
    LEFT JOIN h ON h.contrato_id = ct.id AND h.mes = m.mes
   WHERE (ct.ativo OR h.total IS NOT NULL) {filtro})
SELECT *,
  CASE WHEN tipo = 'mensal_fixo' AND vigente THEN coalesce(valor_mensal, 0) ELSE 0 END AS mensalidade,
  CASE WHEN tipo = 'mensal_fixo' AND franquia_horas IS NOT NULL
       THEN greatest({horas} / 60.0 - franquia_horas, 0) * coalesce(valor_hora, 0)
       WHEN tipo = 'hora_avulsa' THEN {horas} / 60.0 * coalesce(valor_hora, 0)
       ELSE 0 END AS horas_valor,
  CASE WHEN tipo = 'mensal_fixo' AND franquia_horas IS NOT NULL THEN greatest({horas} - franquia_horas * 60, 0)
       WHEN tipo = 'hora_avulsa' THEN {horas} ELSE 0 END AS horas_cobradas
  FROM base"""


def _cobranca(de, ate, cliente, horas="fat_aprov"):
    filtro = " AND cl.id = %s" if cliente else ""
    return db().all(SQL_COBRANCA.format(filtro=filtro, horas=horas), [de, ate] + ([cliente] if cliente else []))


def r_prefaturamento(de, ate, cliente):
    hoje = dt.datetime.now(TZ).date()
    try:
        mes = dt.date.fromisoformat((request.args.get("mes") or hoje.strftime("%Y-%m")) + "-01")
    except ValueError:
        mes = hoje.replace(day=1)
    linhas = [l for l in _cobranca(mes, mes, cliente) if l["mensalidade"] or l["horas_valor"] or l["pendente"]]
    rot = {"mensal_fixo": "Mensal fixo", "banco_horas": "Banco de horas", "hora_avulsa": "Hora avulsa",
           "projeto": "Projeto", "proposta_avulsa": "Proposta avulsa"}
    for l in linhas:
        l["tipo"] = rot.get(l["tipo"], l["tipo"])
        l["total_r"] = round(float(l["mensalidade"]) + float(l["horas_valor"]), 2)
    por_cli = {}
    for l in linhas:
        c = por_cli.setdefault(l["cliente"], {"cliente": l["cliente"], "total_r": 0, "pendente": 0})
        c["total_r"] += l["total_r"]
        c["pendente"] += l["pendente"]
    return {"mes": mes,
            "cards": [("Mês", mes.strftime("%m/%Y"), "texto"),
                      ("Total a faturar (R$)", sum(l["total_r"] for l in linhas), "moeda"),
                      ("Clientes", len(por_cli), "num"),
                      ("Horas pendentes de aprovação", sum(l["pendente"] for l in linhas), "horas")],
            "tabelas": [{"titulo": "Por cliente", "colunas": [("cliente", "Cliente", "texto"),
                                                               ("total_r", "A faturar (R$)", "moeda"),
                                                               ("pendente", "Horas pendentes", "horas")],
                         "linhas": sorted(por_cli.values(), key=lambda x: -x["total_r"])},
                        {"titulo": "Detalhe por contrato", "colunas": [
                            ("cliente", "Cliente", "texto"), ("descricao", "Contrato", "texto"), ("tipo", "Tipo", "texto"),
                            ("mensalidade", "Mensalidade", "moeda"), ("horas_cobradas", "Horas cobradas", "horas"),
                            ("valor_hora", "Valor hora", "moeda"), ("horas_valor", "Valor das horas", "moeda"),
                            ("total_r", "Total (R$)", "moeda"), ("pendente", "Horas pendentes", "horas")],
                         "linhas": linhas}],
            "nota": "Considera só horas faturáveis APROVADAS. Aprove as pendentes antes de faturar. "
                    "Banco de horas é pré-pago; projetos e propostas serão faturados por etapa (Fase 2)."}


def r_rentabilidade(de, ate, cliente):
    linhas = _cobranca(de, ate, cliente, horas="fat")
    por = {}
    for l in linhas:
        receita = float(l["mensalidade"]) + float(l["horas_valor"])
        if l["tipo"] in ("banco_horas", "projeto", "proposta_avulsa"):
            receita += l["fat"] / 60 * float(l["valor_hora"] or 0)   # valor das horas consumidas
        k = (l["cliente"], l["descricao"])
        r = por.setdefault(k, {"cliente": l["cliente"], "contrato": l["descricao"], "receita": 0.0, "custo": 0.0,
                               "minutos": 0})
        r["receita"] += receita
        r["custo"] += float(l["custo"])
        r["minutos"] += l["total"]
    itens = [r for r in por.values() if r["receita"] or r["custo"]]
    for r in itens:
        r["margem"] = round(r["receita"] - r["custo"], 2)
        r["pct"] = round(100 * r["margem"] / r["receita"], 1) if r["receita"] else None
        r["receita"], r["custo"] = round(r["receita"], 2), round(r["custo"], 2)
    clientes = {}
    for r in itens:
        c = clientes.setdefault(r["cliente"], {"cliente": r["cliente"], "receita": 0, "custo": 0, "minutos": 0})
        for k in ("receita", "custo", "minutos"):
            c[k] += r[k]
    for c in clientes.values():
        c["margem"] = round(c["receita"] - c["custo"], 2)
        c["pct"] = round(100 * c["margem"] / c["receita"], 1) if c["receita"] else None
    rec, cus = sum(r["receita"] for r in itens), sum(r["custo"] for r in itens)
    sem_custo = db().val("SELECT count(*) FROM usuarios WHERE perfil <> 'cliente' AND ativo AND custo_hora IS NULL")
    cols = [("receita", "Receita (R$)", "moeda"), ("custo", "Custo (R$)", "moeda"), ("margem", "Margem (R$)", "moeda"),
            ("pct", "Margem %", "pct"), ("minutos", "Horas", "horas")]
    return {"cards": [("Receita estimada", rec, "moeda"), ("Custo dos recursos", cus, "moeda"),
                      ("Margem", rec - cus, "moeda"), ("Margem %", round(100 * (rec - cus) / rec, 1) if rec else None, "pct")],
            "tabelas": [{"titulo": "Por cliente", "colunas": [("cliente", "Cliente", "texto")] + cols,
                         "linhas": sorted(clientes.values(), key=lambda x: -x["margem"])},
                        {"titulo": "Por contrato", "colunas": [("cliente", "Cliente", "texto"), ("contrato", "Contrato", "texto")] + cols,
                         "linhas": sorted(itens, key=lambda x: -x["margem"])}],
            "nota": ("Receita = mensalidades + horas faturáveis × valor hora de venda do contrato (estimativa, inclui horas "
                     "pendentes). Custo = todas as horas × custo/hora do recurso."
                     + (f" Atenção: {sem_custo} recurso(s) sem custo/hora cadastrado." if sem_custo else ""))}


GERADORES = {"sla": r_sla, "chamados": r_chamados, "produtividade": r_produtividade, "contratos": r_contratos,
             "satisfacao": r_satisfacao, "prefaturamento": r_prefaturamento, "rentabilidade": r_rentabilidade}


@bp.get("/")
@requer("relatorios")
def indice():
    itens = [c for c in CATALOGO if pode(c[3])]
    return render_template("equipe/relatorios_indice.html", itens=itens)


@bp.get("/<chave>")
@requer("relatorios")
def relatorio(chave):
    item = next((c for c in CATALOGO if c[0] == chave), None)
    if not item or chave not in GERADORES:
        abort(404)
    if not pode(item[3]):
        abort(403)
    de, ate, cliente = _periodo()
    dados = GERADORES[chave](de, ate, cliente)
    if request.args.get("formato") == "xlsx":
        return _excel_generico(item[1], de, ate, dados)
    clientes = db().all("SELECT id, coalesce(nome_fantasia, razao_social) AS nome FROM clientes ORDER BY 2")
    return render_template("equipe/relatorio.html", item=item, dados=dados, de=de, ate=ate, cliente=cliente,
                           clientes=clientes, chave=chave)


def _excel_generico(titulo, de, ate, dados):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append([titulo]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Período: {de.strftime('%d/%m/%Y')} a {ate.strftime('%d/%m/%Y')}"])
    ws.append([])
    for rot, val, _fmt in dados["cards"]:
        ws.append([rot, (val / 60 if _fmt == "horas" and isinstance(val, (int, float)) else val)])
    if dados.get("nota"):
        ws.append([]); ws.append([dados["nota"]])
    ws.column_dimensions["A"].width = 42
    for t in dados["tabelas"]:
        sh = wb.create_sheet(t["titulo"][:31].replace("/", "-"))
        sh.append([c[1] for c in t["colunas"]])
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79")
        for l in t["linhas"]:
            linha = []
            for k, _rot, fmt in t["colunas"]:
                v = l.get(k)
                if fmt == "horas" and v is not None:
                    v = round(v / 60, 2)
                elif fmt in ("moeda", "num", "pct") and v is not None:
                    v = float(v)
                linha.append(v)
            sh.append(linha)
        for i, c in enumerate(t["colunas"]):
            sh.column_dimensions[sh.cell(row=1, column=i + 1).column_letter].width = max(12, len(c[1]) + 4)
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    nome = f"{titulo.lower().replace(' ', '_')}_{de}_{ate}.xlsx"
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{__import__('urllib.parse').parse.quote(nome)}"})
