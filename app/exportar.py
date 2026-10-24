"""Exportação genérica (Excel e PDF) de telas a partir da estrutura usada nos relatórios.

dados = {"cards": [(rótulo, valor, formato)], "tabelas": [{"titulo", "colunas": [(chave, rótulo, formato)],
         "linhas": [dict]}], "nota": str}
"""
from __future__ import annotations

import datetime as dt
import io
from urllib.parse import quote

from flask import Response, request

from . import pdf
from .util import TZ


def pedido():
    """Formato pedido na URL (?formato=xlsx|pdf) ou None."""
    f = request.args.get("formato")
    return f if f in ("xlsx", "pdf") else None


def _valor_excel(v, fmt):
    if v is None:
        return None
    if fmt == "horas":
        return round(v / 60, 2)
    if fmt == "horas_dec":
        return float(v)
    if fmt in ("moeda", "num", "pct") and not isinstance(v, (int, float)):
        try:
            return float(v)
        except (TypeError, ValueError):
            return v
    if isinstance(v, dt.datetime):
        return v.astimezone(TZ).replace(tzinfo=None) if v.tzinfo else v
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if v.__class__.__name__ == "Decimal":
        return float(v)
    return v


def excel(titulo, subtitulo, dados) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    primeira = True
    for t in dados.get("tabelas") or []:
        ws = wb.active if primeira else wb.create_sheet()
        primeira = False
        ws.title = "".join(ch for ch in t["titulo"] if ch not in "[]:*?/\\")[:31] or "Dados"
        ws.append([titulo]); ws["A1"].font = Font(bold=True, size=13)
        ws.append([subtitulo])
        ws.append([])
        ws.append([c[1] + (" (h)" if c[2] == "horas" else "") for c in t["colunas"]])
        for c in ws[4]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
        for l in t["linhas"]:
            ws.append([_valor_excel(l.get(k), f) for k, _r, f in t["colunas"]])
        if t.get("total"):
            ws.append([_valor_excel(t["total"].get(k), f) if k in t["total"] else None for k, _r, f in t["colunas"]])
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
        for i, (k, rot, f) in enumerate(t["colunas"], start=1):
            letra = get_column_letter(i)
            maior = max([len(str(rot))] + [len(str(l.get(k) or "")) for l in t["linhas"][:300]])
            ws.column_dimensions[letra].width = min(max(10, maior + 2), 60)
            for r in range(5, ws.max_row + 1):
                cel = ws[f"{letra}{r}"]
                if f == "data":
                    cel.number_format = "DD/MM/YYYY"
                elif f in ("horas", "horas_dec", "num"):
                    cel.number_format = "0.00"
                elif f == "moeda":
                    cel.number_format = '"R$" #,##0.00'
        ws.freeze_panes = "A5"
    if dados.get("cards"):
        ws = wb.create_sheet("Resumo", 0)
        ws.append([titulo]); ws["A1"].font = Font(bold=True, size=13)
        ws.append([subtitulo]); ws.append([])
        for rot, val, fmt in dados["cards"]:
            ws.append([rot, _valor_excel(val, fmt)])
        if dados.get("nota"):
            ws.append([]); ws.append([dados["nota"]])
        ws.column_dimensions["A"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def resposta(formato, titulo, subtitulo, dados, nome_base):
    nome = f"{nome_base}_{dt.datetime.now(TZ):%Y-%m-%d}"
    if formato == "pdf":
        return pdf.resposta(nome + ".pdf", pdf.gerar(titulo, subtitulo, dados))
    return Response(excel(titulo, subtitulo, dados),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nome + '.xlsx')}"})
