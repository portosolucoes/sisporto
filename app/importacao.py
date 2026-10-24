"""Importação de clientes, contratos e contatos a partir de planilha Excel."""
from __future__ import annotations

import datetime as dt
import io
import secrets
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from werkzeug.security import generate_password_hash

from .db import db
from .util import cnpj_valido, parse_decimal, so_digitos, senha_temporaria

TIPOS = {
    "mensal fixo": "mensal_fixo", "banco de horas": "banco_horas", "hora avulsa": "hora_avulsa",
    "projeto": "projeto", "proposta avulsa": "proposta_avulsa",
}

COLUNAS = {
    "Clientes": [("Razão social*", 34), ("Nome fantasia", 24), ("CNPJ", 20), ("E-mail do financeiro", 30),
                 ("Telefone", 16), ("Domínio de e-mail", 22), ("Endereço", 40)],
    "Contratos": [("Cliente* (CNPJ ou nome fantasia)", 30), ("Tipo*", 18), ("Descrição*", 30),
                  ("Valor mensal (R$)", 16), ("Franquia de horas/mês", 18), ("Valor da hora (R$)", 16),
                  ("Horas liberadas", 15), ("Valor total (R$)", 16), ("Início*", 13), ("Fim", 13),
                  ("Dia de vencimento", 16), ("Índice de reajuste", 16), ("Mês do reajuste", 15)],
    "Contatos": [("Cliente* (CNPJ ou nome fantasia)", 30), ("Nome*", 26), ("E-mail*", 32), ("Telefone", 16)],
}

EXEMPLOS = {
    "Clientes": ["Alfa Logística Ltda", "Alfa Logística", "11.222.333/0001-81", "financeiro@alfalog.com.br",
                 "(11) 3333-4444", "alfalog.com.br", "Rua Exemplo, 100 - São Paulo/SP"],
    "Contratos": ["Alfa Logística", "Mensal fixo", "Suporte mensal – 10h", 2500, 10, 180, None, None,
                  dt.date(2026, 1, 1), None, 10, "IPCA", 1],
    "Contatos": ["Alfa Logística", "Maria Souza", "maria@alfalog.com.br", "(11) 98888-7777"],
}

INSTRUCOES = [
    "Como preencher",
    "1. Apague as linhas de exemplo (linha 2 de cada aba) antes de importar.",
    "2. Clientes: um por linha. CNPJ é opcional, mas ajuda a ligar contratos e contatos.",
    "3. Domínio de e-mail: e-mails desse domínio enviados ao suporte@ viram chamados do cliente.",
    "4. Contratos: na coluna Cliente use o CNPJ ou o nome fantasia exatamente como na aba Clientes.",
    "5. Tipo: Mensal fixo, Banco de horas, Hora avulsa, Projeto ou Proposta avulsa (lista na célula).",
    "   • Mensal fixo → Valor mensal (franquia e valor da hora excedente são opcionais)",
    "   • Banco de horas → Horas liberadas (e valor da hora, se quiser)",
    "   • Hora avulsa → Valor da hora",
    "   • Projeto / Proposta avulsa → Valor total (Horas liberadas opcional, para controlar o saldo)",
    "6. Datas no formato dd/mm/aaaa.",
    "7. Contatos: pessoas do cliente que terão acesso ao portal. Cada uma recebe uma senha temporária",
    "   mostrada ao final da importação — envie a cada contato.",
    "8. Se houver qualquer erro, nada é gravado: o sistema mostra a linha e o motivo para você corrigir.",
]


def gerar_modelo() -> bytes:
    wb = Workbook()
    inst = wb.active
    inst.title = "Instruções"
    for i, linha in enumerate(INSTRUCOES, 1):
        inst.cell(row=i, column=1, value=linha).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    inst.column_dimensions["A"].width = 100
    cab = PatternFill("solid", fgColor="1F4E79")
    for aba, cols in COLUNAS.items():
        ws = wb.create_sheet(aba)
        for j, (nome, largura) in enumerate(cols, 1):
            c = ws.cell(row=1, column=j, value=nome)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = cab
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[c.column_letter].width = largura
        for j, v in enumerate(EXEMPLOS[aba], 1):
            cel = ws.cell(row=2, column=j, value=v)
            cel.font = Font(italic=True, color="7F7F7F")
            if isinstance(v, dt.date):
                cel.number_format = "DD/MM/YYYY"
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32
        if aba == "Contratos":
            dv = DataValidation(type="list", formula1='"Mensal fixo,Banco de horas,Hora avulsa,Projeto,Proposta avulsa"',
                                allow_blank=False)
            ws.add_data_validation(dv)
            dv.add("B2:B500")
            for col in ("I", "J"):
                for r in range(2, 501):
                    ws[f"{col}{r}"].number_format = "DD/MM/YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
def _txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    return parse_decimal(str(v))


def _int(v):
    n = _num(v)
    return int(n) if n is not None else None


def _data(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"data inválida '{s}' (use dd/mm/aaaa)")


def _linhas(ws, ncols):
    exemplo = EXEMPLOS.get(ws.title, [])[:3]
    for i, row in enumerate(ws.iter_rows(min_row=2, max_col=ncols, values_only=True), start=2):
        if not any(_txt(v) for v in row):
            continue
        if i == 2 and [_txt(v) for v in row[:3]] == [_txt(v) for v in exemplo]:
            continue  # linha de exemplo do modelo esquecida na planilha
        yield i, list(row) + [None] * (ncols - len(row))


def importar(conteudo: bytes):
    """Valida e grava. Devolve (resumo, erros, acessos). Com erros, nada é gravado."""
    try:
        wb = load_workbook(io.BytesIO(conteudo), data_only=True)
    except Exception:
        return None, ["O arquivo não é uma planilha Excel (.xlsx) válida."], []
    faltando = [a for a in COLUNAS if a not in wb.sheetnames]
    if faltando:
        return None, [f"Aba não encontrada: {', '.join(faltando)}. Use o modelo do sistema."], []

    erros, clientes, contratos, contatos = [], [], [], []

    # --- clientes ---------------------------------------------------------
    vistos = set()
    for n, r in _linhas(wb["Clientes"], 7):
        razao, fantasia, cnpj = _txt(r[0]), _txt(r[1]), so_digitos(_txt(r[2]))
        if not razao:
            erros.append(f"Clientes, linha {n}: falta a razão social."); continue
        if cnpj and not cnpj_valido(cnpj):
            erros.append(f"Clientes, linha {n}: CNPJ inválido ({r[2]})."); continue
        chave = cnpj or (fantasia or razao).lower()
        if chave in vistos:
            erros.append(f"Clientes, linha {n}: cliente repetido na planilha."); continue
        vistos.add(chave)
        if cnpj and db().val("SELECT 1 FROM clientes WHERE cnpj = %s", [cnpj]):
            erros.append(f"Clientes, linha {n}: CNPJ já cadastrado no sistema."); continue
        dominio = (_txt(r[5]) or "").lower().lstrip("@") or None
        clientes.append({"linha": n, "razao": razao, "fantasia": fantasia or razao, "cnpj": cnpj,
                         "email_fin": _txt(r[3]), "tel": _txt(r[4]), "dominio": dominio, "end": _txt(r[6])})

    def achar_cliente(ref):
        """Devolve ('novo', índice) ou ('existe', id) ou None."""
        ref = _txt(ref)
        if not ref:
            return None
        dig = so_digitos(ref)
        for i, c in enumerate(clientes):
            if (dig and len(dig) == 14 and c["cnpj"] == dig) or c["fantasia"].lower() == ref.lower() \
                    or c["razao"].lower() == ref.lower():
                return ("novo", i)
        row = None
        if dig and len(dig) == 14:
            row = db().one("SELECT id FROM clientes WHERE cnpj = %s", [dig])
        if not row:
            row = db().one("""SELECT id FROM clientes WHERE lower(nome_fantasia) = lower(%s)
                              OR lower(razao_social) = lower(%s) LIMIT 1""", [ref, ref])
        return ("existe", row["id"]) if row else None

    # --- contratos --------------------------------------------------------
    for n, r in _linhas(wb["Contratos"], 13):
        try:
            cli = achar_cliente(r[0])
            if not cli:
                raise ValueError(f"cliente '{r[0]}' não encontrado na aba Clientes nem no sistema")
            tipo = TIPOS.get((_txt(r[1]) or "").lower())
            if not tipo:
                raise ValueError("tipo inválido (use a lista da célula)")
            desc = _txt(r[2])
            if not desc:
                raise ValueError("falta a descrição")
            v = {"valor_mensal": _num(r[3]), "franquia_horas": _num(r[4]), "valor_hora": _num(r[5]),
                 "horas_pacote": _num(r[6]), "valor_total": _num(r[7]), "data_inicio": _data(r[8]),
                 "data_fim": _data(r[9]), "dia_vencimento": _int(r[10]), "indice_reajuste": _txt(r[11]),
                 "mes_reajuste": _int(r[12])}
            if not v["data_inicio"]:
                raise ValueError("falta a data de início")
            obrig = {"mensal_fixo": ("valor_mensal", "Valor mensal"), "banco_horas": ("horas_pacote", "Horas liberadas"),
                     "hora_avulsa": ("valor_hora", "Valor da hora"), "projeto": ("valor_total", "Valor total"),
                     "proposta_avulsa": ("valor_total", "Valor total")}[tipo]
            if v[obrig[0]] is None:
                raise ValueError(f"para '{_txt(r[1])}' preencha {obrig[1]}")
            if v["dia_vencimento"] is not None and not 1 <= v["dia_vencimento"] <= 28:
                raise ValueError("dia de vencimento deve ser de 1 a 28")
            if v["mes_reajuste"] is not None and not 1 <= v["mes_reajuste"] <= 12:
                raise ValueError("mês do reajuste deve ser de 1 a 12")
            contratos.append({"linha": n, "cli": cli, "tipo": tipo, "descricao": desc, **v})
        except ValueError as e:
            erros.append(f"Contratos, linha {n}: {e}.")

    # --- contatos ---------------------------------------------------------
    emails = set()
    for n, r in _linhas(wb["Contatos"], 4):
        cli = achar_cliente(r[0])
        nome, email = _txt(r[1]), (_txt(r[2]) or "").lower()
        if not cli:
            erros.append(f"Contatos, linha {n}: cliente '{r[0]}' não encontrado."); continue
        if not nome or "@" not in email:
            erros.append(f"Contatos, linha {n}: informe nome e e-mail."); continue
        chave = (cli, email)
        ja = cli[0] == "existe" and db().val(
            "SELECT 1 FROM usuarios WHERE email = %s AND cliente_id = %s AND perfil = 'cliente'", [email, cli[1]])
        if chave in emails or ja:
            erros.append(f"Contatos, linha {n}: {email} já é usuário deste cliente."); continue
        emails.add(chave)
        contatos.append({"cli": cli, "nome": nome, "email": email, "tel": _txt(r[3])})

    if erros:
        return None, erros, []
    if not (clientes or contratos or contatos):
        return None, ["A planilha está vazia."], []

    # --- gravação ---------------------------------------------------------
    ids = []
    for c in clientes:
        ids.append(db().val("""INSERT INTO clientes (razao_social, nome_fantasia, cnpj, email_financeiro, telefone,
                                                     dominio_email, endereco)
                               VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING id""",
                            [c["razao"], c["fantasia"], c["cnpj"], c["email_fin"], c["tel"], c["dominio"], c["end"]]))

    def cid(cli):
        return ids[cli[1]] if cli[0] == "novo" else cli[1]

    campos = ["tipo", "descricao", "valor_mensal", "franquia_horas", "valor_hora", "horas_pacote", "valor_total",
              "data_inicio", "data_fim", "dia_vencimento", "indice_reajuste", "mes_reajuste"]
    for k in contratos:
        db().run(f"INSERT INTO contratos (cliente_id, {', '.join(campos)}) VALUES (%s, {', '.join(['%s'] * len(campos))})",
                 [cid(k["cli"])] + [k[c] for c in campos])
    acessos = []
    for p in contatos:
        senha = senha_temporaria()
        db().run("""INSERT INTO usuarios (nome, email, perfil, cliente_id, telefone, senha_hash)
                    VALUES (%s,%s,'cliente',%s,%s,%s)""",
                 [p["nome"], p["email"], cid(p["cli"]), p["tel"], generate_password_hash(senha)])
        acessos.append({"nome": p["nome"], "email": p["email"], "senha": senha})
    resumo = {"clientes": len(clientes), "contratos": len(contratos), "contatos": len(contatos)}
    return resumo, [], acessos
