"""Portal do cliente: painel de horas, chamados e conversa com a equipe."""
from __future__ import annotations

import datetime as dt

from flask import Blueprint, flash, g, redirect, render_template, request, url_for

from . import anexos as anx
from . import chamados as ch
from .auth import perfis
from .db import db
from .email_integracao import notificar
from .util import TZ, rotulo

bp = Blueprint("portal", __name__)
cliente = perfis("cliente")


def _int(v):
    try:
        return int(v) if v not in (None, "") else None
    except ValueError:
        return None


def _cid():
    return g.usuario["cliente_id"]


@bp.get("/")
@cliente
def painel():
    cid = _cid()
    hoje = dt.datetime.now(TZ).date()
    mes = hoje.replace(day=1)
    empresa = db().one("SELECT * FROM clientes WHERE id = %s", [cid])
    horas_mes = db().one("""SELECT coalesce(sum(minutos_total),0) AS total,
                                   coalesce(sum(minutos_faturaveis_aprovados),0) AS aprovadas,
                                   coalesce(sum(chamados_atendidos),0) AS chamados
                              FROM vw_horas_mes WHERE cliente_id = %s AND mes = %s""", [cid, mes])
    historico = db().all("""SELECT mes, sum(minutos_total) AS minutos FROM vw_horas_mes
                             WHERE cliente_id = %s AND mes > %s GROUP BY mes ORDER BY mes""",
                         [cid, (mes - dt.timedelta(days=170)).replace(day=1)])
    bancos = db().all("""SELECT * FROM vw_saldo_banco_horas s JOIN contratos c ON c.id = s.contrato_id
                          WHERE s.cliente_id = %s AND c.ativo""", [cid])
    franquias = db().all("""SELECT f.*, c.descricao FROM contratos c
                              LEFT JOIN vw_franquia_mes f ON f.contrato_id = c.id AND f.mes = %s
                             WHERE c.cliente_id = %s AND c.ativo AND c.tipo = 'mensal_fixo'
                               AND c.franquia_horas IS NOT NULL""", [mes, cid])
    por_tipo = db().all("""SELECT coalesce(ts.nome, 'Sem tipo') AS tipo, sum(a.minutos) AS minutos
                             FROM apontamentos a JOIN contratos ct ON ct.id = a.contrato_id
                             LEFT JOIN tipos_servico ts ON ts.id = a.tipo_servico_id
                            WHERE ct.cliente_id = %s AND a.status <> 'rejeitado' AND a.data >= %s AND a.data < %s
                            GROUP BY 1 ORDER BY 2 DESC""", [cid, mes, (mes + dt.timedelta(days=32)).replace(day=1)])
    abertos = db().all("""SELECT * FROM vw_chamados_sla WHERE cliente_id = %s
                           AND status IN ('novo','em_atendimento','aguardando_cliente')
                           ORDER BY criado_em DESC""", [cid])
    return render_template("portal/painel.html", empresa=empresa, horas_mes=horas_mes, historico=historico,
                           bancos=bancos, franquias=franquias, abertos=abertos, mes=mes, por_tipo=por_tipo)


@bp.get("/chamados")
@cliente
def chamados():
    situacao = request.args.get("situacao", "abertos")
    sql = "SELECT * FROM vw_chamados_sla WHERE cliente_id = %s"
    if situacao == "abertos":
        sql += " AND status IN ('novo','em_atendimento','aguardando_cliente')"
    elif situacao == "encerrados":
        sql += " AND status IN ('resolvido','fechado')"
    lista = db().all(sql + " ORDER BY criado_em DESC LIMIT 200", [_cid()])
    return render_template("portal/chamados.html", lista=lista, situacao=situacao)


@bp.route("/chamados/novo", methods=["GET", "POST"])
@cliente
def chamado_novo():
    if request.method == "POST":
        try:
            anexo = anx.validar(request.files.get("anexo"))
            novo = ch.criar(cliente_id=_cid(), solicitante_id=g.usuario["id"], origem="portal",
                            titulo=request.form.get("titulo"), descricao=request.form.get("descricao"),
                            prioridade=request.form.get("prioridade", "media"),
                            categoria=request.form.get("categoria"))
            if anexo:
                anx.salvar(novo["id"], None, anexo, g.usuario["id"])
            notificar(novo["id"], "aberto")
            flash(f"Chamado #{novo['numero']} aberto. Você receberá as atualizações por aqui e por e-mail.", "ok")
            return redirect(url_for("portal.chamado", numero=novo["numero"]))
        except ValueError as e:
            flash(str(e), "erro")
    return render_template("portal/chamado_novo.html", form=request.form, prioridades=ch.PRIORIDADES)


@bp.get("/chamados/<int:numero>")
@cliente
def chamado(numero):
    c = ch.buscar(numero, cliente_id=_cid())
    tempo = ch.linha_do_tempo(c["id"], incluir_internas=False)
    horas = db().all("""SELECT data, minutos, descricao FROM apontamentos
                         WHERE chamado_id = %s AND status <> 'rejeitado' AND faturavel ORDER BY data""", [c["id"]])
    return render_template("portal/chamado.html", c=c, tempo=tempo, horas=horas,
                           anexos=anx.listar(c["id"], incluir_internos=False))


@bp.post("/chamados/<int:numero>/responder")
@cliente
def chamado_responder(numero):
    c = ch.buscar(numero, cliente_id=_cid())
    if c["status"] == "fechado":
        flash("Este chamado está fechado. Abra um novo, se precisar.", "erro")
    else:
        try:
            anexo = anx.validar(request.files.get("anexo"))
            iid = ch.responder(c["id"], g.usuario["id"], request.form.get("texto"))
            if anexo:
                anx.salvar(c["id"], iid, anexo, g.usuario["id"])
            if c["status"] == "resolvido" and request.form.get("reabrir") == "1":
                db().run("UPDATE chamados SET status = 'em_atendimento' WHERE id = %s", [c["id"]])
            flash("Mensagem enviada.", "ok")
        except ValueError as e:
            flash(str(e), "erro")
    return redirect(url_for("portal.chamado", numero=numero) + "#conversa")


@bp.post("/chamados/<int:numero>/avaliar")
@cliente
def chamado_avaliar(numero):
    c = ch.buscar(numero, cliente_id=_cid())
    nota = request.form.get("nota", "")
    if c["status"] in ("resolvido", "fechado") and nota in ("1", "2", "3", "4", "5"):
        db().run("""UPDATE chamados SET avaliacao = %s, avaliacao_comentario = %s,
                           status = 'fechado' WHERE id = %s""",
                 [int(nota), (request.form.get("comentario") or "").strip() or None, c["id"]])
        flash("Obrigado pela avaliação!", "ok")
    return redirect(url_for("portal.chamado", numero=numero))


def _filtros_horas():
    hoje = dt.datetime.now(TZ).date()
    a = request.args
    try:
        if a.get("mes"):
            de = dt.date.fromisoformat(a["mes"] + "-01")
            ate = (de + dt.timedelta(days=32)).replace(day=1) - dt.timedelta(days=1)
        else:
            de = dt.date.fromisoformat(a["de"]) if a.get("de") else hoje.replace(day=1)
            ate = dt.date.fromisoformat(a["ate"]) if a.get("ate") else hoje
    except ValueError:
        de, ate = hoje.replace(day=1), hoje
    if ate < de:
        de, ate = ate, de
    f = {"de": de, "ate": ate, "contrato": _int(a.get("contrato")), "tipo": _int(a.get("tipo")),
         "chamado": _int((a.get("chamado") or "").lstrip("#")), "visao": a.get("visao") or "analitica"}
    where = ["ct.cliente_id = %s", "a.data BETWEEN %s AND %s", "a.status <> 'rejeitado'"]
    p = [_cid(), de, ate]
    for k, col in (("contrato", "ct.id"), ("tipo", "a.tipo_servico_id"), ("chamado", "c.numero")):
        if f[k]:
            where.append(f"{col} = %s"); p.append(f[k])
    return f, " AND ".join(where), p


BASE_HORAS = """FROM apontamentos a
          JOIN contratos ct ON ct.id = a.contrato_id
          LEFT JOIN tipos_servico ts ON ts.id = a.tipo_servico_id
          LEFT JOIN chamados c ON c.id = a.chamado_id
          JOIN usuarios u ON u.id = a.usuario_id"""


def _dados_horas():
    f, where, p = _filtros_horas()
    linhas = db().all(f"""
        SELECT a.data, a.minutos, a.descricao, a.faturavel, a.status, c.numero, c.titulo, c.status::text AS ch_status,
               u.nome AS recurso, ct.descricao AS contrato, coalesce(ts.nome, 'Sem tipo') AS tipo
          {BASE_HORAS} WHERE {where} ORDER BY c.numero NULLS LAST, a.data, a.id""", p)
    por_chamado = db().all(f"""
        SELECT c.numero, coalesce(c.titulo, 'Atividades do contrato (sem chamado)') AS titulo,
               c.status::text AS status, ct.descricao AS contrato, count(*) AS lancamentos,
               min(a.data) AS primeiro, max(a.data) AS ultimo, sum(a.minutos) AS minutos,
               coalesce(sum(a.minutos) FILTER (WHERE a.faturavel), 0) AS faturavel
          {BASE_HORAS} WHERE {where} GROUP BY 1, 2, 3, 4 ORDER BY 1 NULLS LAST""", p)
    por_tipo = db().all(f"""SELECT coalesce(ts.nome, 'Sem tipo') AS nome, sum(a.minutos) AS minutos,
                                   coalesce(sum(a.minutos) FILTER (WHERE a.faturavel), 0) AS faturavel
                              {BASE_HORAS} WHERE {where} GROUP BY 1 ORDER BY 2 DESC""", p)
    total = sum(l["minutos"] for l in linhas)
    fat = sum(l["minutos"] for l in linhas if l["faturavel"])
    return f, linhas, por_chamado, por_tipo, total, fat


def _estrutura_horas(f, linhas, por_chamado, por_tipo, total, fat):
    """Mesma estrutura dos relatórios da equipe (usada no PDF)."""
    for l in linhas:
        l["chamado"] = f"#{l['numero']}" if l["numero"] else "—"
        l["fat_txt"] = "Sim" if l["faturavel"] else "Cortesia"
    for c in por_chamado:
        c["chamado"] = f"#{c['numero']}" if c["numero"] else "—"
        c["status"] = rotulo(c["status"], "status") if c["status"] else "—"
    tabelas = [
        {"titulo": "Resumo por chamado", "larguras": [7, 34, 10, 20, 7, 8, 8],
         "colunas": [("chamado", "Chamado", "texto"), ("titulo", "Título", "texto"), ("status", "Situação", "texto"),
                     ("contrato", "Contrato", "texto"), ("lancamentos", "Lanç.", "num"),
                     ("minutos", "Horas", "horas"), ("faturavel", "Faturáveis", "horas")],
         "linhas": por_chamado, "total": {"titulo": "Total", "minutos": total, "faturavel": fat}},
        {"titulo": "Por tipo de serviço", "colunas": [("nome", "Tipo de serviço", "texto"),
                                                     ("minutos", "Horas", "horas"), ("faturavel", "Faturáveis", "horas")],
         "linhas": por_tipo},
    ]
    if f["visao"] == "analitica":
        tabelas.append({"titulo": "Detalhe dos apontamentos", "larguras": [8, 7, 20, 12, 35, 13, 7, 7],
                        "colunas": [("data", "Data", "data"), ("chamado", "Chamado", "texto"),
                                    ("titulo", "Título", "texto"), ("tipo", "Tipo", "texto"),
                                    ("descricao", "Atividade", "texto"), ("recurso", "Recurso", "texto"),
                                    ("minutos", "Horas", "horas"), ("fat_txt", "Faturável", "texto")],
                        "linhas": linhas, "total": {"descricao": "Total do período", "minutos": total}})
    return {"cards": [("Total no período", total, "horas"), ("Faturáveis", fat, "horas"),
                      ("Chamados atendidos", sum(1 for c in por_chamado if c["numero"]), "num"),
                      ("Lançamentos", len(linhas), "num")],
            "tabelas": tabelas,
            "nota": "Horas em revisão podem ser ajustadas pela Porto Soluções até a aprovação."}


@bp.get("/horas")
@cliente
def horas():
    f, linhas, por_chamado, por_tipo, total, fat = _dados_horas()
    formato = request.args.get("formato")
    nome_cli = db().val("SELECT coalesce(nome_fantasia, razao_social) FROM clientes WHERE id = %s", [_cid()])
    if formato == "pdf":
        from . import pdf
        dados = _estrutura_horas(f, linhas, por_chamado, por_tipo, total, fat)
        sub = f"{nome_cli} · Período: {f['de']:%d/%m/%Y} a {f['ate']:%d/%m/%Y}"
        return pdf.resposta(f"horas_{f['de']}_{f['ate']}.pdf", pdf.gerar("Relatório de horas", sub, dados))
    if formato == "xlsx":
        return _horas_excel(f, nome_cli, linhas, por_chamado, por_tipo)
    meses = db().all("""SELECT DISTINCT to_char(mes, 'YYYY-MM') AS m FROM vw_horas_mes
                         WHERE cliente_id = %s ORDER BY 1 DESC""", [_cid()])
    contratos = db().all("SELECT id, descricao FROM contratos WHERE cliente_id = %s ORDER BY ativo DESC, descricao",
                         [_cid()])
    tipos = db().all("SELECT id, nome FROM tipos_servico ORDER BY ordem, nome")
    detalhe = {}
    for l in linhas:
        detalhe.setdefault(l["numero"], []).append(l)
    return render_template("portal/horas.html", f=f, linhas=linhas, por_chamado=por_chamado, por_tipo=por_tipo,
                           detalhe=detalhe, total=total, faturaveis=fat, meses=[m["m"] for m in meses],
                           contratos=contratos, tipos=tipos)


def _horas_excel(f, nome_cli, linhas, por_chamado, por_tipo):
    import io

    from flask import Response
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def cabecalho(ws):
        for c in ws[ws.max_row]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")

    wb = Workbook()
    ws = wb.active
    ws.title = "Por chamado"
    ws.append([f"{nome_cli} – horas de {f['de']:%d/%m/%Y} a {f['ate']:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Chamado", "Título", "Situação", "Contrato", "Lançamentos", "Primeiro", "Último", "Horas", "Faturáveis"])
    cabecalho(ws)
    for c in por_chamado:
        ws.append([c["numero"], c["titulo"], rotulo(c["status"], "status") if c["status"] else None, c["contrato"],
                   c["lancamentos"], c["primeiro"], c["ultimo"], round(c["minutos"] / 60, 2),
                   round(c["faturavel"] / 60, 2)])
    ws.append(["", "Total", "", "", sum(c["lancamentos"] for c in por_chamado), "", "",
               round(sum(c["minutos"] for c in por_chamado) / 60, 2),
               round(sum(c["faturavel"] for c in por_chamado) / 60, 2)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHI", (10, 44, 16, 28, 12, 11, 11, 9, 11)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Apontamentos")
    ws2.append(["Data", "Chamado", "Título", "Contrato", "Tipo de serviço", "Atividade", "Recurso", "Horas",
                "Faturável", "Situação"])
    cabecalho(ws2)
    for l in linhas:
        ws2.append([l["data"], l["numero"], l["titulo"], l["contrato"], l["tipo"], l["descricao"], l["recurso"],
                    round(l["minutos"] / 60, 2), "Sim" if l["faturavel"] else "Cortesia",
                    "Aprovada" if l["status"] == "aprovado" else "Em revisão"])
    for col, w in zip("ABCDEFGHIJ", (11, 9, 34, 24, 18, 50, 20, 8, 10, 12)):
        ws2.column_dimensions[col].width = w
    for r in range(2, ws2.max_row + 1):
        ws2[f"A{r}"].number_format = "DD/MM/YYYY"
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    ws3 = wb.create_sheet("Por tipo de serviço")
    ws3.append(["Tipo de serviço", "Horas", "Faturáveis"])
    cabecalho(ws3)
    for t in por_tipo:
        ws3.append([t["nome"], round(t["minutos"] / 60, 2), round(t["faturavel"] / 60, 2)])
    ws3.column_dimensions["A"].width = 30
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=horas_{f['de']}_{f['ate']}.xlsx"})
