"""Administração: consultas SQL (somente leitura), dicionário de dados e dashboard gerencial."""
from __future__ import annotations

import datetime as dt
import io
import re

from flask import Blueprint, Response, flash, g, redirect, render_template, request, url_for

from . import dicionario
from .acessos import requer
from .auth import admin
from .db import DatabaseError, db
from .util import TZ

bp = Blueprint("adm", __name__)

LIMITE_TELA = 500
LIMITE_EXCEL = 50000
COLUNAS_SIGILOSAS = {"senha_hash", "totp_segredo", "token_hash"}
INICIO_PERMITIDO = re.compile(r"^\s*(select|with|explain|table|values|show)\b", re.I)


def _hoje():
    return dt.datetime.now(TZ).date()


# =============================================================================
# Consultas SQL
# =============================================================================
def validar_sql(sql: str) -> str:
    sql = (sql or "").strip()
    while sql.endswith(";"):
        sql = sql[:-1].rstrip()
    if not sql:
        raise ValueError("Escreva a consulta.")
    if ";" in sql:
        raise ValueError("Execute uma consulta por vez (sem ';' no meio).")
    if not INICIO_PERMITIDO.match(sql):
        raise ValueError("Somente consultas de leitura: SELECT, WITH, EXPLAIN, TABLE, VALUES ou SHOW.")
    if len(sql) > 20000:
        raise ValueError("Consulta muito longa.")
    return sql


def executar(sql: str, limite: int):
    """Roda em transação SOMENTE LEITURA com tempo máximo de 20 s. Devolve (colunas, linhas, truncou)."""
    con = db()
    con.commit()  # grava o que a requisição já fez e abre uma transação nova só para a consulta
    try:
        con.run("SET TRANSACTION READ ONLY")
        con.run("SET LOCAL statement_timeout = '20s'")
        sql_exec = sql
        if re.match(r"^\s*(select|with|table|values)\b", sql, re.I):
            sql_exec = f"SELECT * FROM ({sql}) AS consulta LIMIT {limite + 1}"
        colunas, linhas = con.bruto(sql_exec)
    finally:
        con.rollback()  # nada do que a consulta fizer fica gravado
    truncou = len(linhas) > limite
    linhas = linhas[:limite]
    sig = [i for i, c in enumerate(colunas) if c.lower() in COLUNAS_SIGILOSAS]
    if sig:
        linhas = [tuple("•••" if i in sig and v is not None else v for i, v in enumerate(l)) for l in linhas]
    return colunas, linhas, truncou


def _log(sql, linhas=None, erro=None):
    db().run("INSERT INTO consultas_log (usuario_id, sql, linhas, erro) VALUES (%s,%s,%s,%s)",
             [g.usuario["id"], sql[:20000], linhas, (erro or "")[:2000] or None])


@bp.route("/admin/consultas", methods=["GET", "POST"])
@admin
def consultas():
    sql = request.values.get("sql") or ""
    resultado = erro = None
    if request.method == "POST":
        acao = request.form.get("acao", "executar")
        if acao == "salvar":
            nome = (request.form.get("nome") or "").strip()
            try:
                validar_sql(sql)
                if not nome:
                    raise ValueError("Dê um nome para salvar a consulta.")
                db().run("""INSERT INTO consultas_salvas (nome, sql, criado_por) VALUES (%s,%s,%s)
                            ON CONFLICT (nome) DO UPDATE SET sql = EXCLUDED.sql""", [nome, sql.strip(), g.usuario["id"]])
                flash(f"Consulta “{nome}” salva.", "ok")
            except ValueError as e:
                flash(str(e), "erro")
        elif acao == "excluir":
            db().run("DELETE FROM consultas_salvas WHERE id = %s", [int(request.form.get("id") or 0)])
            flash("Consulta excluída.", "ok")
            return redirect(url_for("adm.consultas"))
        else:
            excel = acao == "excel"
            try:
                limpo = validar_sql(sql)
                colunas, linhas, truncou = executar(limpo, LIMITE_EXCEL if excel else LIMITE_TELA)
                _log(limpo, len(linhas))
                if excel:
                    return _excel(colunas, linhas, limpo)
                resultado = {"colunas": colunas, "linhas": linhas, "truncou": truncou}
            except ValueError as e:
                erro = str(e)
                if sql.strip():
                    _log(sql.strip(), erro=f"bloqueada: {erro}")
            except DatabaseError as e:
                erro = str(e).strip().split("\n")[0]
                _log(sql.strip(), erro=erro)
    elif request.args.get("salva"):
        s = db().one("SELECT sql FROM consultas_salvas WHERE id = %s", [int(request.args["salva"] or 0)])
        sql = s["sql"] if s else ""
    salvas = db().all("SELECT id, nome, sql FROM consultas_salvas ORDER BY nome")
    historico = db().all("""SELECT l.sql, l.linhas, l.erro, l.executado_em, u.nome FROM consultas_log l
                              JOIN usuarios u ON u.id = l.usuario_id ORDER BY l.id DESC LIMIT 15""")
    tabelas = db().all("""SELECT c.relname AS nome, CASE c.relkind WHEN 'v' THEN 'visão' ELSE 'tabela' END AS tipo,
                                 (SELECT string_agg(a.attname, ', ' ORDER BY a.attnum) FROM pg_attribute a
                                   WHERE a.attrelid = c.oid AND a.attnum > 0 AND NOT a.attisdropped) AS colunas
                            FROM pg_class c JOIN pg_namespace n ON n.oid = c.relnamespace
                           WHERE n.nspname = 'public' AND c.relkind IN ('r','v') ORDER BY c.relkind, c.relname""")
    return render_template("admin/consultas.html", sql=sql, resultado=resultado, erro=erro, salvas=salvas,
                           historico=historico, tabelas=tabelas, limite=LIMITE_TELA)


def _excel(colunas, linhas, sql):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Consulta"
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    def cel(v):
        if isinstance(v, dt.datetime):
            return v.astimezone(TZ).replace(tzinfo=None) if v.tzinfo else v
        if v is None or isinstance(v, (bool, int, float, str, dt.date)):
            return v
        if v.__class__.__name__ == "Decimal":
            return float(v)
        return str(v)
    for l in linhas:
        ws.append([cel(v) for v in l])
    ws.freeze_panes = "A2"
    sh = wb.create_sheet("SQL")
    sh["A1"] = sql
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=consulta_{_hoje()}.xlsx"})


# =============================================================================
# Dicionário de dados
# =============================================================================
@bp.get("/admin/banco")
@admin
def banco():
    d = dicionario.ler()
    if request.args.get("formato") == "pdf":
        from . import pdf
        return pdf.resposta(f"sisporto_dicionario_de_dados_{_hoje()}.pdf",
                            pdf.gerar("SisPorto – Dicionário de dados", f"Estrutura do banco em {_hoje():%d/%m/%Y}",
                                      dicionario.estrutura_pdf(d)))
    return render_template("admin/banco.html", d=d)


# =============================================================================
# Dashboard gerencial
# =============================================================================
@bp.get("/dashboard")
@requer("dashboard")
def dashboard():
    hoje = _hoje()
    try:
        mes = dt.date.fromisoformat((request.args.get("mes") or "") + "-01")
    except ValueError:
        mes = hoje.replace(day=1)
    fim = (mes + dt.timedelta(days=32)).replace(day=1) - dt.timedelta(days=1)
    ch = db().one("""
        SELECT count(*) FILTER (WHERE criado_em::date BETWEEN %s AND %s) AS abertos,
               count(*) FILTER (WHERE resolvido_em::date BETWEEN %s AND %s) AS resolvidos,
               count(*) FILTER (WHERE resolvido_em::date BETWEEN %s AND %s
                                  AND (prazo_solucao IS NULL OR resolvido_em <= prazo_solucao)) AS no_prazo,
               count(*) FILTER (WHERE status IN ('novo','em_atendimento','aguardando_cliente')) AS backlog,
               count(*) FILTER (WHERE situacao_sla = 'estourado'
                                  AND status IN ('novo','em_atendimento','aguardando_cliente')) AS estourados,
               round(avg(avaliacao) FILTER (WHERE resolvido_em::date BETWEEN %s AND %s), 2) AS csat
          FROM vw_chamados_sla""", [mes, fim] * 4)
    hr = db().one("""SELECT coalesce(sum(minutos), 0) AS total,
                            coalesce(sum(minutos) FILTER (WHERE faturavel), 0) AS faturavel,
                            coalesce(sum(minutos) FILTER (WHERE status = 'pendente'), 0) AS pendente
                       FROM apontamentos WHERE status <> 'rejeitado' AND data BETWEEN %s AND %s""", [mes, fim])
    mrr = db().val("SELECT coalesce(sum(valor_mensal), 0) FROM contratos WHERE ativo AND valor_mensal IS NOT NULL")
    nps = db().one("""SELECT count(*) FILTER (WHERE nota IS NOT NULL) AS respostas,
                             round(100.0 * (count(*) FILTER (WHERE nota >= 9) - count(*) FILTER (WHERE nota <= 6))
                                   / nullif(count(*) FILTER (WHERE nota IS NOT NULL), 0)) AS nps
                        FROM nps_respostas WHERE mes = %s""", [mes])
    try:
        from .relatorios import r_prefaturamento
        prefat = next((v for r, v, _f in r_prefaturamento(mes, fim, None)["cards"] if "faturar" in r), None)
    except Exception:  # noqa: BLE001 - indicador opcional
        db().rollback()
        prefat = None
    serie = db().all("""SELECT to_char(date_trunc('month', data), 'MM/YYYY') AS mes, sum(minutos) AS minutos,
                               coalesce(sum(minutos) FILTER (WHERE faturavel), 0) AS faturavel
                          FROM apontamentos WHERE status <> 'rejeitado' AND data >= %s AND data <= %s
                         GROUP BY date_trunc('month', data) ORDER BY date_trunc('month', data)""",
                     [(mes - dt.timedelta(days=150)).replace(day=1), fim])
    top = db().all("""SELECT coalesce(cl.nome_fantasia, cl.razao_social) AS nome, sum(a.minutos) AS minutos
                        FROM apontamentos a JOIN contratos ct ON ct.id = a.contrato_id
                        JOIN clientes cl ON cl.id = ct.cliente_id
                       WHERE a.status <> 'rejeitado' AND a.data BETWEEN %s AND %s
                       GROUP BY 1 ORDER BY 2 DESC LIMIT 8""", [mes, fim])
    por_tipo = db().all("""SELECT coalesce(ts.nome, 'Sem tipo') AS nome, sum(a.minutos) AS minutos
                             FROM apontamentos a LEFT JOIN tipos_servico ts ON ts.id = a.tipo_servico_id
                            WHERE a.status <> 'rejeitado' AND a.data BETWEEN %s AND %s
                            GROUP BY 1 ORDER BY 2 DESC""", [mes, fim])
    alertas = db().all("""
        SELECT ct.id, coalesce(cl.nome_fantasia, cl.razao_social) AS cliente, ct.descricao, s.horas_saldo,
               s.pct_consumido, ct.data_fim,
               CASE WHEN s.pct_consumido >= 80 THEN 'saldo' ELSE 'vencimento' END AS motivo
          FROM contratos ct JOIN clientes cl ON cl.id = ct.cliente_id
          LEFT JOIN vw_saldo_banco_horas s ON s.contrato_id = ct.id
         WHERE ct.ativo AND (s.pct_consumido >= 80 OR ct.data_fim BETWEEN %s AND %s)
         ORDER BY s.pct_consumido DESC NULLS LAST, ct.data_fim""", [hoje, hoje + dt.timedelta(days=60)])
    meses = [(hoje.replace(day=1) - dt.timedelta(days=31 * i)).replace(day=1) for i in range(12)]
    return render_template("admin/dashboard.html", mes=mes, ch=ch, hr=hr, mrr=mrr, nps=nps, prefat=prefat,
                           serie=serie, top=top, por_tipo=por_tipo, alertas=alertas, meses=meses)
